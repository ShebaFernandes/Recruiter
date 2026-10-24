"""Create a varied fictional Gemini resume benchmark without external requests."""

# ruff: noqa: E501

from __future__ import annotations

import hashlib
import json
import sys
from dataclasses import asdict, dataclass
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "scripts"))

from resume_accuracy_benchmark import pdf_bytes  # noqa: E402
from talent.resume_understanding import (  # noqa: E402
    interpretation_text,
    request_input_token_reservation,
)
from talent.services import parse_resume  # noqa: E402

OUTPUT = ROOT / ".local-benchmark" / "gemini-varied"
MODEL = "gemini-3.5-flash-lite"
MAX_OUTPUT_TOKENS = 3000
INPUT_PRICE = Decimal("0.30")
OUTPUT_PRICE = Decimal("2.50")


@dataclass(frozen=True)
class Role:
    company: str
    title: str
    start: str | None
    end: str | None
    current: bool
    achievement: str


@dataclass(frozen=True)
class Case:
    slug: str
    kind: str
    title: str
    summary: str
    skills: tuple[str, ...]
    roles: tuple[Role, ...]
    lines: tuple[str, ...]
    layout: str = "paragraphs"


def cases() -> tuple[Case, ...]:
    return (
        Case("promotion-gap", "pdf", "Staff Platform Engineer", "Builds reliable delivery platforms.", ("Python", "Django", "AWS", "Kubernetes", "PostgreSQL"), (
            Role("Blue Acorn Robotics", "Platform Engineer", "March 2017", "August 2019", False, "Reduced deployment recovery time by 31%."),
            Role("Blue Acorn Robotics", "Senior Platform Engineer", "September 2019", "January 2021", False, "Cut build queue time from 18 minutes to 7 minutes."),
            Role("Copper Finch Systems", "Staff Platform Engineer", "April 2021", None, True, "Improved release success from 91% to 98%."),
        ), ("Fictional Resume 01", "Title: Staff Platform Engineer", "Summary: Builds reliable delivery platforms.", "Skills: Python, Django, AWS, Kubernetes, PostgreSQL", "PROFESSIONAL EXPERIENCE", "Blue Acorn Robotics", "Platform Engineer", "March 2017 - August 2019", "Reduced deployment recovery time by 31%.", "Senior Platform Engineer", "September 2019 - January 2021", "Cut build queue time from 18 minutes to 7 minutes.", "Copper Finch Systems", "Staff Platform Engineer", "April 2021 - Present", "Improved release success from 91% to 98%.")),
        Case("overlapping-advisory", "docx", "Principal Data Engineer", "Designs observable data systems.", ("Python", "Kafka", "SQL", "AWS"), (
            Role("Harbour Metric Labs", "Senior Data Engineer", "June 2018", "May 2024", False, "Raised pipeline completion from 94% to 99.4%."),
            Role("Ivory Lake Analytics", "Data Platform Advisor", "September 2022", "February 2023", False, "Planned migration of 26 production feeds."),
            Role("Juniper Signal Works", "Principal Data Engineer", "June 2024", None, True, "Reduced late events by 42%."),
        ), ("Fictional Resume 02", "Title: Principal Data Engineer", "Summary: Designs observable data systems.", "Skills: Python, Kafka, SQL, AWS", "PROFESSIONAL EXPERIENCE", "Harbour Metric Labs", "Senior Data Engineer", "June 2018 - May 2024", "Raised pipeline completion from 94% to 99.4%.", "Ivory Lake Analytics", "Data Platform Advisor", "September 2022 - February 2023", "Planned migration of 26 production feeds.", "Juniper Signal Works", "Principal Data Engineer", "June 2024 - Present", "Reduced late events by 42%.")),
        Case("compact-year-promotions", "pdf", "Engineering Manager", "Leads product engineering teams.", ("Java", "Spring", "Leadership"), (
            Role("Kiteframe Systems", "Software Engineer", "2012", "2014", False, "Automated 40 regression checks."),
            Role("Kiteframe Systems", "Senior Software Engineer", "2014", "2015", False, "Reduced release defects by 24%."),
            Role("Kiteframe Systems", "Engineering Manager", "2016", "2020", False, "Grew the engineering group from 5 to 14."),
        ), ("Fictional Resume 03", "Title: Engineering Manager", "Summary: Leads product engineering teams.", "Skills: Java, Spring, Leadership", "Kiteframe Systems | 2012-2020", "Software Engineer | 2012-14", "Automated 40 regression checks.", "Senior Software Engineer | 2014-15", "Reduced release defects by 24%.", "Engineering Manager | 2016-20", "Grew the engineering group from 5 to 14.")),
        Case("distant-employer-heading", "docx", "Product Engineering Lead", "Builds accessible SaaS products.", ("TypeScript", "React", "Node.js"), (
            Role("Lantern Grove Software", "Frontend Engineer", "January 2019", "December 2020", False, "Improved accessibility audit score to 98%."),
            Role("Lantern Grove Software", "Senior Frontend Engineer", "January 2021", "March 2023", False, "Cut page-load time by 36%."),
            Role("Lantern Grove Software", "Product Engineering Lead", "April 2023", None, True, "Delivered four releases without rollback."),
        ), ("Fictional Resume 04", "Title: Product Engineering Lead", "Summary: Builds accessible SaaS products.", "Skills: TypeScript, React, Node.js", "PROFESSIONAL EXPERIENCE", "Lantern Grove Software | January 2019 - Present", "Selected progression at the same employer", "Frontend Engineer", "January 2019 - December 2020", "Improved accessibility audit score to 98%.", "Senior Frontend Engineer", "January 2021 - March 2023", "Cut page-load time by 36%.", "Product Engineering Lead", "April 2023 - Present", "Delivered four releases without rollback.")),
        Case("rehire-same-employer", "pdf", "Security Engineer", "Secures cloud services.", ("AWS", "Python", "Terraform"), (
            Role("Mosaic Defence Labs", "Security Analyst", "February 2016", "July 2018", False, "Closed 87 high-risk findings."),
            Role("Northwind Commerce", "Cloud Security Engineer", "October 2018", "May 2021", False, "Reduced exposed services by 63%."),
            Role("Mosaic Defence Labs", "Senior Security Engineer", "August 2021", None, True, "Cut incident triage time by 41%."),
        ), ("Fictional Resume 05", "Title: Security Engineer", "Summary: Secures cloud services.", "Skills: AWS, Python, Terraform", "Mosaic Defence Labs", "Security Analyst | February 2016 - July 2018", "Closed 87 high-risk findings.", "Northwind Commerce", "Cloud Security Engineer | October 2018 - May 2021", "Reduced exposed services by 63%.", "Mosaic Defence Labs", "Senior Security Engineer | August 2021 - Present", "Cut incident triage time by 41%.")),
        Case("table-format", "docx", "Operations Manager", "Improves service operations.", ("Operations", "Excel", "SQL"), (
            Role("Oakline Services", "Operations Analyst", "May 2015", "June 2018", False, "Reduced weekly reporting time by 11 hours."),
            Role("Pinebridge Logistics", "Operations Manager", "July 2018", None, True, "Improved on-time delivery from 88% to 96%."),
        ), ("Fictional Resume 06", "Title: Operations Manager", "Summary: Improves service operations.", "Skills: Operations, Excel, SQL", "Oakline Services | Operations Analyst | May 2015 - June 2018 | Reduced weekly reporting time by 11 hours.", "Pinebridge Logistics | Operations Manager | July 2018 - Present | Improved on-time delivery from 88% to 96%."), "table"),
        Case("explicit-career-gap", "pdf", "UX Research Lead", "Turns research into product decisions.", ("Research", "Figma", "Analytics"), (
            Role("Quartz Product Studio", "UX Researcher", "March 2017", "November 2020", False, "Completed 62 moderated research sessions."),
            Role("Riverstone Health Tech", "UX Research Lead", "September 2021", None, True, "Raised research adoption across six teams."),
        ), ("Fictional Resume 07", "Title: UX Research Lead", "Summary: Turns research into product decisions.", "Skills: Research, Figma, Analytics", "Quartz Product Studio", "UX Researcher | March 2017 - November 2020", "Completed 62 moderated research sessions.", "Career break | December 2020 - August 2021 | Family care", "Riverstone Health Tech", "UX Research Lead | September 2021 - Present", "Raised research adoption across six teams.")),
        Case("contract-and-permanent", "docx", "Finance Systems Lead", "Modernises finance platforms.", ("SAP", "SQL", "Power BI"), (
            Role("Silver Birch Finance", "Systems Analyst", "January 2016", "December 2019", False, "Reconciled 12 legacy data sources."),
            Role("Tamarind Advisory", "ERP Consultant", "April 2019", "October 2019", False, "Delivered a six-month migration plan."),
            Role("Umbra Payments", "Finance Systems Lead", "February 2020", None, True, "Reduced month-end close by three days."),
        ), ("Fictional Resume 08", "Title: Finance Systems Lead", "Summary: Modernises finance platforms.", "Skills: SAP, SQL, Power BI", "Silver Birch Finance", "Systems Analyst | January 2016 - December 2019", "Reconciled 12 legacy data sources.", "Tamarind Advisory", "ERP Consultant (contract) | April 2019 - October 2019", "Delivered a six-month migration plan.", "Umbra Payments", "Finance Systems Lead | February 2020 - Present", "Reduced month-end close by three days.")),
        Case("unknown-start-date", "pdf", "Technical Writer", "Creates clear developer documentation.", ("Documentation", "Git", "APIs"), (
            Role("Violet API Works", "Technical Writer", None, "June 2020", False, "Documented 48 public endpoints."),
            Role("Willow Developer Tools", "Senior Technical Writer", "August 2020", None, True, "Reduced support tickets by 29%."),
        ), ("Fictional Resume 09", "Title: Technical Writer", "Summary: Creates clear developer documentation.", "Skills: Documentation, Git, APIs", "Violet API Works", "Technical Writer | Start date not stated - June 2020", "Documented 48 public endpoints.", "Willow Developer Tools", "Senior Technical Writer | August 2020 - Present", "Reduced support tickets by 29%.")),
        Case("similar-titles", "docx", "Quality Engineering Lead", "Builds reliable test systems.", ("Python", "Playwright", "CI/CD"), (
            Role("Xenon Mobile", "Quality Engineer", "January 2018", "March 2020", False, "Automated 73 mobile checks."),
            Role("Yellow Finch Apps", "Senior Quality Engineer", "April 2020", "December 2022", False, "Reduced escaped defects by 33%."),
            Role("Zenith Product Labs", "Quality Engineering Lead", "January 2023", None, True, "Cut regression time from two days to four hours."),
        ), ("Fictional Resume 10", "Title: Quality Engineering Lead", "Summary: Builds reliable test systems.", "Skills: Python, Playwright, CI/CD", "Xenon Mobile | Quality Engineer | January 2018 - March 2020 | Automated 73 mobile checks.", "Yellow Finch Apps | Senior Quality Engineer | April 2020 - December 2022 | Reduced escaped defects by 33%.", "Zenith Product Labs | Quality Engineering Lead | January 2023 - Present | Cut regression time from two days to four hours."), "table"),
        Case("six-role-history", "pdf", "Director of Engineering", "Leads growing engineering organisations.", ("Leadership", "Go", "Python", "AWS"), tuple(
            Role(f"Fictional Company {index}", title, f"January {2010 + index * 2}", f"December {2011 + index * 2}" if index < 5 else None, index == 5, f"Delivered measurable programme {index + 1} on schedule.")
            for index, title in enumerate(("Engineer", "Senior Engineer", "Tech Lead", "Engineering Manager", "Head of Engineering", "Director of Engineering"))
        ), tuple(["Fictional Resume 11", "Title: Director of Engineering", "Summary: Leads growing engineering organisations.", "Skills: Leadership, Go, Python, AWS", "PROFESSIONAL EXPERIENCE"] + [item for index, title in enumerate(("Engineer", "Senior Engineer", "Tech Lead", "Engineering Manager", "Head of Engineering", "Director of Engineering")) for item in (f"Fictional Company {index}", f"{title} | January {2010 + index * 2} - {'Present' if index == 5 else f'December {2011 + index * 2}'}", f"Delivered measurable programme {index + 1} on schedule.")])),
        Case("short-current-role", "docx", "Machine Learning Engineer", "Deploys measured machine-learning services.", ("Python", "PyTorch", "Docker"), (
            Role("Aster Model Works", "Machine Learning Engineer", "April 2024", None, True, "Reduced inference latency by 27%."),
        ), ("Fictional Resume 12", "Title: Machine Learning Engineer", "Summary: Deploys measured machine-learning services.", "Skills: Python, PyTorch, Docker", "Aster Model Works", "Machine Learning Engineer | April 2024 - Present", "Reduced inference latency by 27%.")),
    )


def build(case: Case) -> bytes:
    if case.kind == "pdf":
        return pdf_bytes(list(case.lines))
    document = Document()
    if case.layout == "table":
        for line in case.lines[:4]:
            document.add_paragraph(line)
        table = document.add_table(rows=0, cols=4)
        for line in case.lines[4:]:
            values = [value.strip() for value in line.split("|", 3)]
            row = table.add_row().cells
            for cell, value in zip(row, values, strict=True):
                cell.text = value
    else:
        for line in case.lines:
            document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    results = []
    for case in cases():
        path = OUTPUT / f"{case.slug}.{case.kind}"
        content = build(case)
        path.write_bytes(content)
        with path.open("rb") as source:
            text = interpretation_text(parse_resume(source, path.name))
        required = [case.title, case.summary, *case.skills]
        for role in case.roles:
            required.extend((role.company, role.title, role.achievement))
        missing = [value for value in required if value.lower() not in text.lower()]
        if missing:
            raise RuntimeError(f"{case.slug} lost expected evidence: {missing}")
        reserved_input = request_input_token_reservation(text)
        reservation = (
            Decimal(reserved_input) * INPUT_PRICE
            + Decimal(MAX_OUTPUT_TOKENS) * OUTPUT_PRICE
        ) / Decimal("1000000")
        results.append({
            "slug": case.slug,
            "path": str(path),
            "sha256": hashlib.sha256(content).hexdigest(),
            "expected": {
                "professional_title": case.title,
                "summary": case.summary,
                "skills": case.skills,
                "work_experiences": [asdict(role) for role in case.roles],
            },
            "document_characters": len(text),
            "all_expected_evidence_present": True,
            "reserved_input_tokens": reserved_input,
            "max_output_tokens": MAX_OUTPUT_TOKENS,
            "paid_equivalent_reservation_usd": str(reservation.quantize(Decimal("0.000001"))),
        })
    manifest = {
        "mode": "offline-only",
        "external_requests": 0,
        "model": MODEL,
        "document_count": len(results),
        "documents": results,
    }
    (OUTPUT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps({
        "mode": manifest["mode"],
        "external_requests": 0,
        "document_count": len(results),
        "all_expected_evidence_present": all(
            result["all_expected_evidence_present"] for result in results
        ),
        "maximum_reservation_usd": max(
            result["paid_equivalent_reservation_usd"] for result in results
        ),
        "manifest": str(OUTPUT / "manifest.json"),
    }, indent=2))


if __name__ == "__main__":
    main()
