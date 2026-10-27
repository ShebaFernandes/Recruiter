"""Synthetic, offline employment-history benchmark. No candidate data or API calls."""

# ruff: noqa: E501

import argparse
import json
import subprocess
import sys
import time
import types
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))


@dataclass(frozen=True)
class Role:
    company: str
    title: str
    start: str | None = None
    end: str | None = None
    current: bool = False


@dataclass(frozen=True)
class Case:
    name: str
    lines: list[str]
    roles: list[Role]
    kind: str
    table: bool = False
    columns: bool = False
    image_only: bool = False


def cases():
    return [
        Case("chronological_pdf", ["WORK EXPERIENCE", "Engineer | Alpha Labs | Jan 2020 - Dec 2021", "Senior Engineer | Beta Labs | Jan 2022 - Present"], [Role("Alpha Labs", "Engineer", "2020-01-01", "2021-12-01"), Role("Beta Labs", "Senior Engineer", "2022-01-01", current=True)], "pdf"),
        Case("chronological_docx", ["EMPLOYMENT HISTORY", "Developer at Cedar Works", "Jun 2017 - May 2019", "Lead Developer at Delta Works", "Jun 2019 - Dec 2022"], [Role("Cedar Works", "Developer", "2017-06-01", "2019-05-01"), Role("Delta Works", "Lead Developer", "2019-06-01", "2022-12-01")], "docx"),
        Case("promotions_pdf", ["PROFESSIONAL EXPERIENCE", "Evergreen Group Jan 2018 - Present", "Director of Engineering Jan 2023 - Present", "Engineering Manager Jan 2020 - Dec 2022", "Software Engineer Jan 2018 - Dec 2019"], [Role("Evergreen Group", "Director of Engineering", "2023-01-01", current=True), Role("Evergreen Group", "Engineering Manager", "2020-01-01", "2022-12-01"), Role("Evergreen Group", "Software Engineer", "2018-01-01", "2019-12-01")], "pdf"),
        Case("promotions_docx", ["WORK HISTORY", "Fable Systems 2016 - Present", "VP Engineering 2022 - Present", "Senior Engineer 2019 - 2021", "Engineer 2016 - 2018"], [Role("Fable Systems", "VP Engineering", "2022-01-01", current=True), Role("Fable Systems", "Senior Engineer", "2019-01-01", "2021-01-01"), Role("Fable Systems", "Engineer", "2016-01-01", "2018-01-01")], "docx"),
        Case("overlap_pdf", ["EXPERIENCE", "Platform Engineer | Gamma Co | Jan 2020 - Present", "Advisor | Harbor Studio | Jun 2021 - Jun 2022"], [Role("Gamma Co", "Platform Engineer", "2020-01-01", current=True), Role("Harbor Studio", "Advisor", "2021-06-01", "2022-06-01")], "pdf"),
        Case("year_only_docx", ["PROFESSIONAL EXPERIENCE", "Analyst | Indigo | 2014 - 2017", "Senior Analyst | Juniper | 2018 - 2020"], [Role("Indigo", "Analyst", "2014-01-01", "2017-01-01"), Role("Juniper", "Senior Analyst", "2018-01-01", "2020-01-01")], "docx"),
        Case("unknown_dates_pdf", ["WORK HISTORY", "Product Manager at Kite Labs"], [Role("Kite Labs", "Product Manager")], "pdf"),
        Case("intern_contract_docx", ["EXPERIENCE", "Software Intern | Lumen | Jan 2019 - Jun 2019 | Internship", "Consultant | Mosaic | Jul 2019 - Dec 2020 | Contract"], [Role("Lumen", "Software Intern", "2019-01-01", "2019-06-01"), Role("Mosaic", "Consultant", "2019-07-01", "2020-12-01")], "docx"),
        Case("docx_table", ["WORK HISTORY", "North Labs", "Senior Engineer", "2020-02", "Present", "Old Works", "Software Developer", "06/2017", "01-2020"], [Role("North Labs", "Senior Engineer", "2020-02-01", current=True), Role("Old Works", "Software Developer", "2017-06-01", "2020-01-01")], "docx", table=True),
        Case("mixed_months_pdf", ["EMPLOYMENT HISTORY", "Engineer | Pine Co | 2018/03 - 04/2020", "Lead Engineer | Quartz Co | 05-2020 - 2023-09"], [Role("Pine Co", "Engineer", "2018-03-01", "2020-04-01"), Role("Quartz Co", "Lead Engineer", "2020-05-01", "2023-09-01")], "pdf"),
        Case("separate_company_role_docx", ["PROFESSIONAL EXPERIENCE", "River Systems", "Staff Engineer Apr 2021 - Present", "Stone Systems", "Engineer Mar 2018 - Mar 2021"], [Role("River Systems", "Staff Engineer", "2021-04-01", current=True), Role("Stone Systems", "Engineer", "2018-03-01", "2021-03-01")], "docx"),
        Case("current_keyword_pdf", ["EXPERIENCE", "Security Engineer | Timber | Feb 2022 - Current"], [Role("Timber", "Security Engineer", "2022-02-01", current=True)], "pdf"),
        Case("now_keyword_docx", ["WORK EXPERIENCE", "Data Scientist | Umber AI | 2021 - Now"], [Role("Umber AI", "Data Scientist", "2021-01-01", current=True)], "docx"),
        Case("undated_plus_dated_pdf", ["WORK HISTORY", "Designer at Violet Studio", "Design Lead | Willow Studio | Mar 2020 - Dec 2023"], [Role("Violet Studio", "Designer"), Role("Willow Studio", "Design Lead", "2020-03-01", "2023-12-01")], "pdf"),
        Case("same_title_companies_docx", ["EXPERIENCE", "Engineer | Xenon | 2016 - 2018", "Engineer | Yellow | 2018 - 2020"], [Role("Xenon", "Engineer", "2016-01-01", "2018-01-01"), Role("Yellow", "Engineer", "2018-01-01", "2020-01-01")], "docx"),
        Case("duplicate_role_pdf", ["WORK EXPERIENCE", "Engineer | Zenith | Jan 2020 - Dec 2021", "Engineer | Zenith | Jan 2020 - Dec 2021"], [Role("Zenith", "Engineer", "2020-01-01", "2021-12-01")], "pdf"),
        Case("long_history_docx", ["PROFESSIONAL EXPERIENCE", *[f"Engineer | Company {i} | {2008 + i} - {2009 + i}" for i in range(8)]], [Role(f"Company {i}", "Engineer", f"{2008 + i}-01-01", f"{2009 + i}-01-01") for i in range(8)], "docx"),
        Case("career_history_pdf", ["CAREER HISTORY", "Developer | Amber | Jan 2019 - Jan 2020"], [Role("Amber", "Developer", "2019-01-01", "2020-01-01")], "pdf"),
        Case("multicolumn_pdf", ["Name: Synthetic Person", "Skills: Python", "PROFESSIONAL EXPERIENCE", "Senior Engineer | Beacon | Jan 2021 - Present", "Engineer | Arc | Jan 2018 - Dec 2020"], [Role("Beacon", "Senior Engineer", "2021-01-01", current=True), Role("Arc", "Engineer", "2018-01-01", "2020-12-01")], "pdf", columns=True),
        Case("spaced_employer_date_pdf", ["PROFESSIONAL EXPERIENCE", "Harbor Systems                    Jan 2020 - Present", "Senior Software Engineer                    Bengaluru, India", "Built fictional systems.", "SKILLS", "Python"], [Role("Harbor Systems", "Senior Software Engineer", "2020-01-01", current=True)], "pdf"),
        Case("spaced_intern_promotion_docx", ["CAREER HISTORY", "Blue Lantern Labs                    Jan 2018 - Jun 2018", "Engineering Intern                    Pune, India", "Built a fictional prototype.", "Copper Systems                    Jan 2020 - Jun 2023", "Software Engineer                    Pune, India", "Built fictional services.", "Copper Systems                    Jul 2023 - Present", "Senior Software Engineer                    Pune, India", "Led fictional delivery.", "SKILLS", "Python"], [Role("Blue Lantern Labs", "Engineering Intern", "2018-01-01", "2018-06-01"), Role("Copper Systems", "Software Engineer", "2020-01-01", "2023-06-01"), Role("Copper Systems", "Senior Software Engineer", "2023-07-01", current=True)], "docx"),
        Case("fresher_projects_docx", ["Name: Fictional Graduate", "Fresher seeking first role", "PROJECTS", "Built an engineer scheduling prototype in 2025.", "EDUCATION", "Example University"], [], "docx"),
        Case("unrecognized_selected_experience_pdf", ["SELECTED EXPERIENCE", "Developer | Marigold Labs | Jan 2022 - Dec 2024", "SKILLS", "Python"], [Role("Marigold Labs", "Developer", "2022-01-01", "2024-12-01")], "pdf"),
        Case("image_only_pdf", [], [], "pdf", image_only=True),
    ]


def pdf_bytes(lines, columns=False, image_only=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if image_only:
        output = BytesIO()
        writer.write(output)
        return output.getvalue()
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    commands = ["BT", "/F1 10 Tf"]
    if columns:
        left, right = lines[:2], lines[2:]
        for index, line in enumerate(left):
            commands.extend([f"1 0 0 1 50 {740 - index * 18} Tm", f"({line}) Tj"])
        for index, line in enumerate(right):
            commands.extend([f"1 0 0 1 310 {740 - index * 18} Tm", f"({line}) Tj"])
    else:
        for index, line in enumerate(lines):
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.extend([f"1 0 0 1 55 {740 - index * 18} Tm", f"({escaped}) Tj"])
    commands.append("ET")
    stream = DecodedStreamObject()
    stream.set_data("\n".join(commands).encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def document(case):
    if case.kind == "pdf":
        return pdf_bytes(case.lines, case.columns, case.image_only), f"{case.name}.pdf"
    doc = Document()
    if case.table:
        table = doc.add_table(rows=1, cols=4)
        table.rows[0].cells[0].text = case.lines[0]
        for values in (case.lines[1:5], case.lines[5:9]):
            row = table.add_row().cells
            for cell, value in zip(row, values, strict=True):
                cell.text = value
    else:
        for line in case.lines:
            doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    return output.getvalue(), f"{case.name}.docx"


def parser(use_head):
    if not use_head:
        from talent.services import parse_resume
        return parse_resume
    source = subprocess.run(
        ["git", "show", "HEAD:backend/talent/services.py"],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    module = types.ModuleType("legacy_resume_services")
    exec(compile(source, "legacy_resume_services.py", "exec"), module.__dict__)
    return module.parse_resume


def normalized(value):
    return " ".join((value or "").casefold().split())


def run(use_head=False):
    parse_resume = parser(use_head)
    totals = {"employers": 0, "roles": 0, "dates": 0, "employer_hits": 0, "role_hits": 0, "date_hits": 0, "unsupported": 0, "duplicates": 0, "processing_errors": 0, "extracted": 0, "matched": 0}
    durations = []
    case_results = []
    for case in cases():
        content, filename = document(case)
        started = time.perf_counter()
        try:
            result = parse_resume(BytesIO(content), filename)
            rows = result["work_experiences"]
        except Exception:
            rows = []
            totals["processing_errors"] += 1
        durations.append((time.perf_counter() - started) * 1000)
        expected_pairs = {(normalized(role.company), normalized(role.title)) for role in case.roles}
        actual_pairs = [(normalized(row.get("company")), normalized(row.get("role"))) for row in rows]
        totals["employers"] += len(case.roles)
        totals["roles"] += len(case.roles)
        totals["extracted"] += len(rows)
        case_hits = 0
        for role in case.roles:
            if any(company == normalized(role.company) for company, _ in actual_pairs):
                totals["employer_hits"] += 1
            pair = (normalized(role.company), normalized(role.title))
            matching = [row for row, actual in zip(rows, actual_pairs, strict=True) if actual == pair]
            if matching:
                totals["role_hits"] += 1
                totals["matched"] += 1
                case_hits += 1
                row = matching[0]
                for key, expected in (("start_date", role.start), ("end_date", role.end)):
                    if expected is not None:
                        totals["dates"] += 1
                        totals["date_hits"] += row.get(key) == expected
                if role.current:
                    totals["dates"] += 1
                    totals["date_hits"] += bool(row.get("is_current") or "present" in (row.get("date_label") or "").casefold())
        totals["unsupported"] += sum(pair not in expected_pairs for pair in actual_pairs)
        totals["duplicates"] += len(actual_pairs) - len(set(actual_pairs))
        case_results.append({"name": case.name, "kind": case.kind, "expected_roles": len(case.roles), "extracted_roles": len(rows), "matched_roles": case_hits, "false_positive_roles": sum(pair not in expected_pairs for pair in actual_pairs)})
    return {
        "fixture_count": len(cases()),
        "employer_recall": round(totals["employer_hits"] / totals["employers"], 4),
        "employment_role_recall": round(totals["role_hits"] / totals["roles"], 4),
        "employment_role_precision": round(totals["matched"] / totals["extracted"], 4),
        "expected_roles": totals["roles"],
        "extracted_roles": totals["extracted"],
        "date_accuracy": round(totals["date_hits"] / totals["dates"], 4),
        "unsupported_rows": totals["unsupported"],
        "duplicate_rows": totals["duplicates"],
        "processing_errors": totals["processing_errors"],
        "mean_processing_ms": round(sum(durations) / len(durations), 2),
        "max_processing_ms": round(max(durations), 2),
        "cases": case_results,
    }


if __name__ == "__main__":
    arguments = argparse.ArgumentParser()
    arguments.add_argument("--baseline-head", action="store_true")
    args = arguments.parse_args()
    print(json.dumps(run(args.baseline_head), indent=2))
