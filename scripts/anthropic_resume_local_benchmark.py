"""Prepare two fictional resumes and audit local text/rule extraction.

This command is deliberately offline. It has no Anthropic client and cannot make an
external request. The generated files and report live under .local-benchmark/, which is
ignored by Git and can be deleted after the benchmark.
"""

from __future__ import annotations

import hashlib
import json
import sys
import time
from dataclasses import asdict, dataclass
from datetime import date
from io import BytesIO
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT / "scripts"))

from resume_accuracy_benchmark import pdf_bytes  # noqa: E402
from talent.resume_understanding import interpretation_text  # noqa: E402
from talent.services import parse_resume  # noqa: E402


@dataclass(frozen=True)
class ExpectedRole:
    company: str
    role: str
    start_date: str
    end_date: str | None
    is_current: bool
    achievement: str


@dataclass(frozen=True)
class BenchmarkCase:
    slug: str
    filename: str
    professional_title: str
    summary: str
    skills: tuple[str, ...]
    roles: tuple[ExpectedRole, ...]
    lines: tuple[str, ...]
    kind: str


CASES = (
    BenchmarkCase(
        slug="fictional-promotion-gap",
        filename="fictional-promotion-gap.pdf",
        professional_title="Staff Platform Engineer",
        summary=(
            "Platform engineer focused on reliable developer tooling and measurable "
            "delivery improvements."
        ),
        skills=("Python", "Django", "AWS", "Kubernetes", "PostgreSQL"),
        roles=(
            ExpectedRole(
                "Blue Acorn Robotics",
                "Platform Engineer",
                "2017-03-01",
                "2019-08-01",
                False,
                "Reduced deployment recovery time by 31%.",
            ),
            ExpectedRole(
                "Blue Acorn Robotics",
                "Senior Platform Engineer",
                "2019-09-01",
                "2021-01-01",
                False,
                "Cut build queue time from 18 minutes to 7 minutes.",
            ),
            ExpectedRole(
                "Copper Finch Systems",
                "Staff Platform Engineer",
                "2021-04-01",
                None,
                True,
                "Improved release success from 91% to 98%.",
            ),
        ),
        lines=(
            "Fictional Benchmark Candidate One",
            "Title: Staff Platform Engineer",
            "Summary: Platform engineer focused on reliable developer tooling and "
            "measurable delivery improvements.",
            "Skills: Python, Django, AWS, Kubernetes, PostgreSQL",
            "PROFESSIONAL EXPERIENCE",
            "Blue Acorn Robotics",
            "Platform Engineer",
            "Start: March 2017",
            "End: August 2019",
            "Reduced deployment recovery time by 31%.",
            "Blue Acorn Robotics",
            "Senior Platform Engineer",
            "Start: September 2019",
            "End: January 2021",
            "Cut build queue time from 18 minutes to 7 minutes.",
            "Copper Finch Systems",
            "Staff Platform Engineer",
            "Start: April 2021",
            "End: Present",
            "Improved release success from 91% to 98%.",
            "EDUCATION",
            "Bachelor of Engineering, Fictional Lakes Institute",
        ),
        kind="pdf",
    ),
    BenchmarkCase(
        slug="fictional-table-overlap",
        filename="fictional-table-overlap.docx",
        professional_title="Senior Data Engineer",
        summary="Data engineer building observable batch and streaming systems.",
        skills=("Python", "Kafka", "AWS", "PostgreSQL"),
        roles=(
            ExpectedRole(
                "Juniper Metric Labs",
                "Data Engineer",
                "2018-06-01",
                "2020-11-01",
                False,
                "Raised daily pipeline completion from 94% to 99.4%.",
            ),
            ExpectedRole(
                "Kestrel Data Cooperative",
                "Senior Data Engineer",
                "2021-02-01",
                "2024-05-01",
                False,
                "Reduced late-arriving events by 42%.",
            ),
            ExpectedRole(
                "Lantern Grove Analytics",
                "Data Platform Advisor",
                "2023-09-01",
                "2024-02-01",
                False,
                "Delivered a migration plan covering 26 production feeds.",
            ),
        ),
        lines=(
            "Fictional Benchmark Candidate Two",
            "Title: Senior Data Engineer",
            "Summary: Data engineer building observable batch and streaming systems.",
            "Skills: Python, Kafka, AWS, PostgreSQL",
            "PROFESSIONAL EXPERIENCE",
            "Employer: Juniper Metric Labs",
            "Position: Data Engineer",
            "Began: June 2018",
            "Left: November 2020",
            "Evidence: Raised daily pipeline completion from 94% to 99.4%.",
            "Employer: Kestrel Data Cooperative",
            "Position: Senior Data Engineer",
            "Began: February 2021",
            "Left: May 2024",
            "Evidence: Reduced late-arriving events by 42%.",
            "Employer: Lantern Grove Analytics",
            "Position: Data Platform Advisor",
            "Began: September 2023",
            "Left: February 2024",
            "Employment type: Contract",
            "Evidence: Delivered a migration plan covering 26 production feeds.",
            "EDUCATION",
            "Master of Science, Fictional Ridge University",
        ),
        kind="docx",
    ),
)


def build_document(case: BenchmarkCase) -> bytes:
    if case.kind == "pdf":
        return pdf_bytes(list(case.lines))
    document = Document()
    for line in case.lines[:5]:
        document.add_paragraph(line)
    table = document.add_table(rows=0, cols=2)
    for line in case.lines[5:-2]:
        label, value = line.split(":", 1)
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value.strip()
    for line in case.lines[-2:]:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def expected_evidence(case: BenchmarkCase) -> list[str]:
    values = [case.professional_title, case.summary, *case.skills]
    for role in case.roles:
        values.extend((role.company, role.role, role.achievement))
    return list(dict.fromkeys(values))


def _fact(value: str | None, evidence: str = "") -> dict:
    return {"value": value, "evidence": evidence}


def _month_label(value: str) -> str:
    return date.fromisoformat(value).strftime("%B %Y")


def mocked_claude_response(case: BenchmarkCase, document_text: str) -> dict:
    """Build a schema-valid, evidence-bound response for worker integration tests."""
    lines = [line for line in document_text.splitlines() if line.strip()]

    def evidence_line(value: str) -> str:
        return next(line for line in lines if value in line)

    cursor = 0
    experiences = []
    for role in case.roles:
        start_index = next(
            index
            for index in range(cursor, len(lines))
            if role.company in lines[index]
        )
        end_index = next(
            index
            for index in range(start_index, len(lines))
            if role.achievement in lines[index]
        )
        block = lines[start_index : end_index + 1]
        cursor = end_index + 1
        achievement_line = evidence_line(role.achievement)
        experiences.append(
            {
                "company": role.company,
                "role": role.role,
                "start": _month_label(role.start_date),
                "end": "Present" if role.is_current else _month_label(role.end_date or ""),
                "employment_type": (
                    "Contract" if any("Contract" in line for line in block) else None
                ),
                "location": None,
                "employer_evidence": block[0],
                "role_evidence": "\n".join(block[1:]),
                "details": role.achievement,
                "achievements": [_fact(role.achievement, achievement_line)],
            }
        )
    return {
        "full_name": _fact(None),
        "location": _fact(None),
        "professional_title": _fact(
            case.professional_title, evidence_line(case.professional_title)
        ),
        "professional_summary": _fact(case.summary, evidence_line(case.summary)),
        "work_experiences": experiences,
        "skills": [_fact(skill, evidence_line(skill)) for skill in case.skills],
        "standout_signal": _fact(
            case.roles[0].achievement, evidence_line(case.roles[0].achievement)
        ),
        "education": [],
    }


def prepare(output_dir: Path) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for case in CASES:
        content = build_document(case)
        path = output_dir / case.filename
        path.write_bytes(content)
        started = time.perf_counter()
        parsed = parse_resume(BytesIO(content), case.filename)
        elapsed_ms = round((time.perf_counter() - started) * 1000, 2)
        supplied = interpretation_text(parsed)
        evidence = expected_evidence(case)
        results.append(
            {
                "slug": case.slug,
                "path": str(path.resolve()),
                "sha256": hashlib.sha256(content).hexdigest(),
                "expected": {
                    "professional_title": case.professional_title,
                    "summary": case.summary,
                    "skills": list(case.skills),
                    "work_experiences": [asdict(role) for role in case.roles],
                },
                "deterministic": {
                    "headline": parsed["headline"],
                    "summary": parsed["summary"],
                    "skills": parsed["skills"],
                    "work_experiences": parsed["work_experiences"],
                    "employment_parse": parsed["_employment_parse"],
                    "processing_ms": elapsed_ms,
                },
                "claude_input_audit": {
                    "character_count": len(supplied),
                    "all_expected_evidence_present": all(value in supplied for value in evidence),
                    "missing_expected_evidence": [
                        value for value in evidence if value not in supplied
                    ],
                    "direct_identifiers_redacted": "Fictional Benchmark Candidate" not in supplied,
                },
            }
        )
    report = {
        "mode": "offline-only",
        "external_requests": 0,
        "synthetic_documents": len(CASES),
        "results": results,
    }
    report_path = output_dir / "preflight.json"
    report_path.write_text(json.dumps(report, indent=2) + "\n")
    return report


if __name__ == "__main__":
    destination = ROOT / ".local-benchmark" / "anthropic-resume"
    print(json.dumps(prepare(destination), indent=2))
