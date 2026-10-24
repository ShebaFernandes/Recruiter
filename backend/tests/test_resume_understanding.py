"""All provider responses and documents are synthetic; network calls are mocked."""

import json
import urllib.error
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from unittest.mock import Mock

import pytest
from django.utils import timezone
from talent.ai_providers import (
    AIProviderAuthenticationError,
    AIProviderInvalidRequest,
    AIProviderMalformedResponse,
    AIProviderNotFound,
    AIProviderPermissionDenied,
    AIProviderRateLimited,
    AIProviderResult,
    AIProviderTimeout,
    AIProviderUnavailable,
)
from talent.resume_understanding import (
    SCHEMA,
    interpretation_text,
    repair_pending_work_experience_dates,
    request_input_token_reservation,
    resume_provider,
    total_experience,
    understand_resume,
    validate_interpretation,
)
from talent.services import parse_resume

TEXT = """Name: Sample Person
Location: Pune
Title: Senior Engineer
Summary: Builds reliable platforms.
Engineer | Example Studio | Jan 2020 - Dec 2021
Built a Python service used by 12 teams.
Senior Engineer | Example Studio | Jan 2022 - Present
Shipped a Python and Kafka platform.
Skills: Python, Kafka, Java
Education: Example University
"""

PROMOTION_TEXT = """Title: AI Engineering Leader
Summary: Builds evidence-bound systems.
Altran (formerly Aricent) Jan 2010 - Apr 2012
Software Engineer — C/C++ telecom systems (SS7 protocol stacks).
CoCubes.com (Acquired by Aon Hewitt, 2016)
Software Engineer | 2012–14
Senior Software Engineer | 2014–15
AVP — Engineering | Jun 2015–Dec 2017
PeopleStrong (A Goldman Sachs company)
Senior Data Scientist & Technical Leader — Founding DS Hire | Dec 2017–Mar 2021
Delivered supported outcome one.

Senior Engineering Manager — Mobile, Jinie (AI), Workgroups | Apr 2021–Mar 2024
Delivered supported outcome two.
Associate Vice President — AI & Experience (Head of AI Engineering) | Apr 2024–Present
Delivered supported outcome three.
Skills: Python, AI, Engineering Leadership
"""

RECALL_TEXT = """Title: Documentation and Research Lead
Summary: Creates evidence-led products and documentation.
Quartz Product Studio
UX Researcher | March 2017 - November 2020
Completed 62 moderated research sessions.
Career break | December 2020 - August 2021 | Family care
Riverstone Health Tech
UX Research Lead | September 2021 - Present
Raised research adoption across six teams.
Violet API Works
Technical Writer | Start date not stated - June 2020
Documented 48 public endpoints.
Willow Developer Tools
Senior Technical Writer | August 2020 - Present
Reduced support tickets by 29%.
Skills: Research, Figma, Analytics, Documentation, Git, APIs
"""


def fact(value, evidence=None):
    return {"value": value, "evidence": evidence or value or ""}


def response():
    return {
        "full_name": fact("Sample Person", "Name: Sample Person"),
        "location": fact("Pune", "Location: Pune"),
        "professional_title": fact("Senior Engineer", "Title: Senior Engineer"),
        "professional_summary": fact(
            "Builds reliable platforms.", "Summary: Builds reliable platforms."
        ),
        "work_experiences": [
            {"company": "Example Studio", "role": "Engineer", "start": "Jan 2020",
             "end": "Dec 2021", "employment_type": None, "location": None,
             "details": "Built a Python service used by 12 teams.",
             "achievements": [fact("Built a Python service used by 12 teams.")],
             "employer_evidence": "Engineer | Example Studio | Jan 2020 - Dec 2021",
             "role_evidence": "Engineer | Example Studio | Jan 2020 - Dec 2021\n"
                              "Built a Python service used by 12 teams."},
            {"company": "Example Studio", "role": "Senior Engineer", "start": "Jan 2022",
             "end": "Present", "employment_type": None, "location": None,
             "details": "Shipped a Python and Kafka platform.",
             "achievements": [fact("Shipped a Python and Kafka platform.")],
             "employer_evidence": "Senior Engineer | Example Studio | Jan 2022 - Present",
             "role_evidence": "Senior Engineer | Example Studio | Jan 2022 - Present\n"
                              "Shipped a Python and Kafka platform."},
        ],
        "skills": [fact("Java"), fact("Kafka"), fact("Python")],
        "standout_signal": fact("Built a Python service used by 12 teams."),
        "education": [fact("Example University")],
    }


def promotion_response():
    employers = {
        "Altran (formerly Aricent)": "Altran (formerly Aricent)",
        "CoCubes.com": "CoCubes.com (Acquired by Aon Hewitt, 2016)",
        "PeopleStrong": "PeopleStrong (A Goldman Sachs company)",
    }
    roles = [
        (
            "Altran (formerly Aricent)",
            "Software Engineer",
            "Jan 2010",
            "Apr 2012",
            "Software Engineer — C/C++ telecom systems (SS7 protocol stacks).",
        ),
        (
            "CoCubes.com",
            "Software Engineer",
            "2012",
            "2014",
            "Software Engineer | 2012–14",
        ),
        (
            "CoCubes.com",
            "Senior Software Engineer",
            "2014",
            "2015",
            "Senior Software Engineer | 2014–15",
        ),
        (
            "CoCubes.com",
            "AVP — Engineering",
            "Jun 2015",
            "Dec 2017",
            "AVP — Engineering | Jun 2015–Dec 2017",
        ),
        (
            "PeopleStrong",
            "Senior Data Scientist & Technical Leader — Founding DS Hire",
            "Dec 2017",
            "Mar 2021",
            "Senior Data Scientist & Technical Leader — Founding DS Hire | "
            "Dec 2017–Mar 2021",
        ),
        (
            "PeopleStrong",
            "Senior Engineering Manager — Mobile, Jinie (AI), Workgroups",
            "Apr 2021",
            "Mar 2024",
            "Senior Engineering Manager — Mobile, Jinie (AI), Workgroups | "
            "Apr 2021–Mar 2024",
        ),
        (
            "PeopleStrong",
            "Associate Vice President — AI & Experience (Head of AI Engineering)",
            "Apr 2024",
            "Present",
            "Associate Vice President — AI & Experience (Head of AI Engineering) | "
            "Apr 2024–Present",
        ),
    ]
    experiences = []
    for company, role, start, end, role_evidence in roles:
        experiences.append(
            {
                "company": company,
                "role": role,
                "start": start,
                "end": end,
                "employment_type": None,
                "location": None,
                "employer_evidence": employers[company],
                "role_evidence": role_evidence,
                "details": "",
                "achievements": [],
            }
        )
    return {
        "full_name": fact(None),
        "location": fact(None),
        "professional_title": fact(
            "AI Engineering Leader", "Title: AI Engineering Leader"
        ),
        "professional_summary": fact(
            "Builds evidence-bound systems.",
            "Summary: Builds evidence-bound systems.",
        ),
        "work_experiences": experiences,
        "skills": [fact("Python", "Skills: Python, AI, Engineering Leadership")],
        "standout_signal": fact(None),
        "education": [],
    }


def work_response(experiences):
    return {
        "full_name": fact(None),
        "location": fact(None),
        "professional_title": fact(None),
        "professional_summary": fact(None),
        "work_experiences": experiences,
        "skills": [],
        "standout_signal": fact(None),
        "education": [],
    }


def recall_response():
    def work(company, role, start, end, role_evidence, achievement):
        return {
            "company": company,
            "role": role,
            "start": start,
            "end": end,
            "employment_type": None,
            "location": None,
            "employer_evidence": company,
            "role_evidence": role_evidence,
            "details": achievement,
            "achievements": [fact(achievement)] if achievement else [],
        }

    return work_response([
        work(
            "Quartz Product Studio",
            "UX Researcher",
            "March 2017",
            "November 2020",
            "UX Researcher | March 2017 - November 2020",
            "Completed 62 moderated research sessions.",
        ),
        work(
            "Quartz Product Studio",
            "Career break",
            "December 2020",
            "August 2021",
            "Career break | December 2020 - August 2021 | Family care",
            "",
        ),
        work(
            "Riverstone Health Tech",
            "UX Research Lead",
            "September 2021",
            "Present",
            "UX Research Lead | September 2021 - Present",
            "Raised research adoption across six teams.",
        ),
        work(
            "Violet API Works",
            "Technical Writer",
            None,
            "June 2020",
            "Technical Writer | Start date not stated - June 2020",
            "Documented 48 public endpoints.",
        ),
        work(
            "Willow Developer Tools",
            "Senior Technical Writer",
            "August 2020",
            "Present",
            "Senior Technical Writer | August 2020 - Present",
            "Reduced support tickets by 29%.",
        ),
    ])


def rules():
    return {"raw_text": TEXT, "skills": ["Java", "Python"], "work_experiences": [],
            "headline": "Engineer"}


def test_resume_text_over_legacy_limit_remains_available_for_ai():
    content = (
        "Name: Fictional Long Resume\n"
        "PROFESSIONAL EXPERIENCE\n"
        "Employer and role evidence.\n"
        + "Synthetic evidence detail. " * 700
    ).encode()
    parsed = parse_resume(BytesIO(content), "fictional-long-resume.txt")
    assert len(parsed["raw_text"]) > 15_000
    assert parsed["_text_truncated"] is False


def test_input_reservation_includes_instructions_schema_and_document():
    reserved = request_input_token_reservation(TEXT)
    document_only_estimate = (len(TEXT.encode("utf-8")) + 1) // 2 + 256
    assert reserved > document_only_estimate


@pytest.fixture
def enabled(settings):
    settings.RESUME_AI_ENABLED = True
    settings.RESUME_AI_EXTERNAL_PROCESSING_APPROVED = True
    settings.RESUME_AI_ALLOW_PAID = False
    settings.RESUME_AI_PROVIDER = "anthropic"
    settings.RESUME_AI_MODEL = "claude-sonnet-5-5"
    settings.RESUME_AI_ROLLOUT_MODE = "consenting"
    settings.LEGAL_RESUME_AI_VERSION = "enter-resume-ai-consent-test-v1"
    settings.LEGAL_RESUME_AI_SHA256 = "1" * 64


def test_valid_evidence_promotions_dedup_and_backend_computation():
    data = response()
    data["work_experiences"].append(deepcopy(data["work_experiences"][0]))
    parsed = validate_interpretation(data, TEXT)
    assert len(parsed["work_experiences"]) == 2
    assert parsed["headline"] == "Senior Engineer"
    assert parsed["summary"] == "Builds reliable platforms."
    assert parsed["current_company"] == "Example Studio"
    assert parsed["strongest_skills"] == ["Python", "Kafka", "Java"]
    assert parsed["likely_directions"] == ["Staff Engineering"]
    assert parsed["standout_signal"] in TEXT
    assert parsed["work_experiences"][0]["start_date"] == "2020-01-01"
    assert parsed["total_experience"] > 5


def test_schema_mismatch_is_rejected():
    data = response()
    data["expected_salary_lpa"] = 50
    with pytest.raises(AIProviderMalformedResponse) as exc_info:
        validate_interpretation(data, TEXT)
    assert exc_info.value.stage == "schema_validation"


def test_seven_promoted_roles_with_distant_employer_headings_are_supported():
    parsed = validate_interpretation(promotion_response(), PROMOTION_TEXT)
    rows = parsed["work_experiences"]
    assert len(rows) == 7
    assert sum(row["company"] == "CoCubes.com" for row in rows) == 3
    assert sum(row["company"] == "PeopleStrong" for row in rows) == 3
    assert rows[0]["date_label"] == "Jan 2010 — Apr 2012"
    assert rows[0]["start_date"] == "2010-01-01"
    assert rows[0]["end_date"] == "2012-04-01"
    assert rows[1]["date_label"] == "2012 — 2014"
    assert rows[1]["start_date"] is None
    assert rows[1]["end_date"] is None
    assert rows[2]["date_label"] == "2014 — 2015"
    assert rows[-1]["is_current"] is True
    assert rows[-1]["end_date"] is None
    assert parsed["_validation_rejections"] == []


def test_pending_date_repair_changes_only_four_evidence_backed_endpoints():
    rows = validate_interpretation(
        promotion_response(), PROMOTION_TEXT
    )["work_experiences"]
    partial = deepcopy(rows)
    partial[0].update(
        start_date=None,
        end_date=None,
        date_label="Unknown — Unknown",
        confidence_status="partial",
    )
    partial[1].update(
        end_date=None,
        date_label="2012 — Unknown",
        confidence_status="partial",
    )
    partial[2].update(
        end_date=None,
        date_label="2014 — Unknown",
        confidence_status="partial",
    )

    repaired, diagnostics = repair_pending_work_experience_dates(
        partial, PROMOTION_TEXT
    )

    assert diagnostics == {
        "changed_indexes": [0, 1, 2],
        "unresolved_indexes": [],
        "validation_rejections": [],
    }
    assert sum(
        endpoint != "Unknown"
        for row in repaired
        for endpoint in row["date_label"].split(" — ")
    ) == 14
    assert repaired[0]["date_label"] == "Jan 2010 — Apr 2012"
    assert repaired[1]["date_label"] == "2012 — 2014"
    assert repaired[2]["date_label"] == "2014 — 2015"
    assert [row["company"] for row in repaired] == [row["company"] for row in partial]
    assert [row["role"] for row in repaired] == [row["role"] for row in partial]


@pytest.mark.parametrize(
    "separator",
    ["-", "–", "—", " - ", " – ", " — "],
)
def test_compact_year_ranges_expand_only_the_supported_end_year(separator):
    text = (
        "Example Studio\n"
        f"Software Engineer | 2012{separator}14\n"
        "Other Studio\n"
    )
    data = work_response([
        {
            "company": "Example Studio",
            "role": "Software Engineer",
            "start": "2012",
            "end": "2014",
            "employment_type": None,
            "location": None,
            "employer_evidence": "Example Studio",
            "role_evidence": f"Software Engineer | 2012{separator}14",
            "details": "",
            "achievements": [],
        }
    ])
    parsed = validate_interpretation(data, text)
    row = parsed["work_experiences"][0]
    assert row["date_label"] == "2012 — 2014"
    assert row["start_date"] is None
    assert row["end_date"] is None
    assert row["confidence_status"] == "supported"
    assert parsed["_validation_rejections"] == []


@pytest.mark.parametrize(
    ("evidence_range", "claimed_end"),
    [("2012–11", "2011"), ("1999–02", "2002"), ("2012–14", "2114")],
)
def test_ambiguous_or_invented_compact_end_years_remain_unresolved(
    evidence_range, claimed_end
):
    text = f"Example Studio\nSoftware Engineer | {evidence_range}\n"
    data = work_response([
        {
            "company": "Example Studio",
            "role": "Software Engineer",
            "start": evidence_range[:4],
            "end": claimed_end,
            "employment_type": None,
            "location": None,
            "employer_evidence": "Example Studio",
            "role_evidence": f"Software Engineer | {evidence_range}",
            "details": "",
            "achievements": [],
        }
    ])
    parsed = validate_interpretation(data, text)
    row = parsed["work_experiences"][0]
    assert row["end_date"] is None
    assert row["date_label"].endswith("— Unknown")
    assert {
        "category": "work_experience.date_evidence",
        "item_index": 0,
        "field": "end",
    } in parsed["_validation_rejections"]


def test_adjacent_employer_heading_dates_support_one_unique_role():
    text = (
        "Altran (formerly Aricent) Jan 2010 — Apr 2012\n"
        "Software Engineer — C/C++ telecom systems.\n"
        "Next Employer May 2012 — Present\n"
        "Senior Engineer\n"
    )
    data = work_response([
        {
            "company": "Altran (formerly Aricent)",
            "role": "Software Engineer",
            "start": "Jan 2010",
            "end": "Apr 2012",
            "employment_type": None,
            "location": None,
            "employer_evidence": "Altran (formerly Aricent)",
            "role_evidence": "Software Engineer — C/C++ telecom systems.",
            "details": "",
            "achievements": [],
        },
        {
            "company": "Next Employer",
            "role": "Senior Engineer",
            "start": "May 2012",
            "end": "Present",
            "employment_type": None,
            "location": None,
            "employer_evidence": "Next Employer",
            "role_evidence": "Senior Engineer",
            "details": "",
            "achievements": [],
        },
    ])
    parsed = validate_interpretation(data, text)
    assert parsed["work_experiences"][0]["date_label"] == "Jan 2010 — Apr 2012"
    assert parsed["work_experiences"][1]["date_label"] == "May 2012 — Present"
    assert parsed["_validation_rejections"] == []


def test_employer_heading_dates_are_not_shared_across_promotions():
    text = (
        "Example Studio Jan 2010 — Dec 2020\n"
        "Software Engineer\n"
        "Engineering Manager\n"
        "Next Employer Jan 2021 — Present\n"
        "Director\n"
    )
    data = work_response([
        {
            "company": "Example Studio",
            "role": role,
            "start": "Jan 2010",
            "end": "Dec 2020",
            "employment_type": None,
            "location": None,
            "employer_evidence": "Example Studio",
            "role_evidence": role,
            "details": "",
            "achievements": [],
        }
        for role in ("Software Engineer", "Engineering Manager")
    ])
    parsed = validate_interpretation(data, text)
    assert len(parsed["work_experiences"]) == 2
    assert all(row["date_label"] == "Unknown — Unknown" for row in parsed["work_experiences"])
    assert sum(
        rejection["category"] == "work_experience.date_evidence"
        for rejection in parsed["_validation_rejections"]
    ) == 4


def test_unrelated_employer_heading_dates_cannot_support_another_role():
    text = (
        "First Employer Jan 2010 — Apr 2012\n"
        "Software Engineer\n"
        "Second Employer May 2012 — Present\n"
        "Engineering Manager\n"
    )
    data = work_response([
        {
            "company": "First Employer",
            "role": "Software Engineer",
            "start": "May 2012",
            "end": "Present",
            "employment_type": None,
            "location": None,
            "employer_evidence": "First Employer",
            "role_evidence": "Software Engineer",
            "details": "",
            "achievements": [],
        }
    ])
    parsed = validate_interpretation(data, text)
    assert parsed["work_experiences"][0]["date_label"] == "Unknown — Unknown"
    assert {
        (rejection.get("field"), rejection["category"])
        for rejection in parsed["_validation_rejections"]
    } == {
        ("start", "work_experience.date_evidence"),
        ("end", "work_experience.current_status"),
    }


def test_unique_employer_heading_dates_preserve_overlapping_positions():
    text = (
        "Main Employer Jan 2020 — Dec 2022\n"
        "Platform Engineer\n"
        "Side Studio Jun 2021 — Jun 2022\n"
        "Independent Advisor\n"
    )
    data = work_response([
        {
            "company": company,
            "role": role,
            "start": start,
            "end": end,
            "employment_type": None,
            "location": None,
            "employer_evidence": company,
            "role_evidence": role,
            "details": "",
            "achievements": [],
        }
        for company, role, start, end in (
            ("Main Employer", "Platform Engineer", "Jan 2020", "Dec 2022"),
            ("Side Studio", "Independent Advisor", "Jun 2021", "Jun 2022"),
        )
    ])
    parsed = validate_interpretation(data, text)
    assert [row["date_label"] for row in parsed["work_experiences"]] == [
        "Jan 2020 — Dec 2022",
        "Jun 2021 — Jun 2022",
    ]
    assert parsed["total_experience"] == 2.9
    assert parsed["_validation_rejections"] == []


def test_wrong_employer_association_and_invented_title_are_rejected_per_role():
    data = promotion_response()
    data["work_experiences"][4].update(
        company="CoCubes.com",
        employer_evidence="CoCubes.com (Acquired by Aon Hewitt, 2016)",
    )
    data["work_experiences"][5]["role"] = "Invented Chief Architect"
    parsed = validate_interpretation(data, PROMOTION_TEXT)
    assert len(parsed["work_experiences"]) == 5
    assert {rejection["category"] for rejection in parsed["_validation_rejections"]} == {
        "work_experience.employer_association",
        "work_experience.role_evidence",
    }


def test_researcher_role_survives_explicit_career_break_without_making_break_a_job():
    parsed = validate_interpretation(recall_response(), RECALL_TEXT)
    rows = parsed["work_experiences"]

    assert [(row["company"], row["role"]) for row in rows[:2]] == [
        ("Quartz Product Studio", "UX Researcher"),
        ("Riverstone Health Tech", "UX Research Lead"),
    ]
    researcher = rows[0]
    assert researcher["start_date"] == "2017-03-01"
    assert researcher["end_date"] == "2020-11-01"
    assert researcher["description"] == "Completed 62 moderated research sessions."
    assert all(row["role"] != "Career break" for row in rows)
    assert {
        "category": "work_experience.non_employment",
        "item_index": 1,
        "field": "role",
    } in parsed["_validation_rejections"]


def test_technical_writer_roles_preserve_unknown_start_and_explicit_current_status():
    parsed = validate_interpretation(recall_response(), RECALL_TEXT)
    rows = {row["role"]: row for row in parsed["work_experiences"]}

    unknown_start = rows["Technical Writer"]
    assert unknown_start["company"] == "Violet API Works"
    assert unknown_start["start_date"] is None
    assert unknown_start["end_date"] == "2020-06-01"
    assert unknown_start["date_label"] == "Unknown — June 2020"
    assert unknown_start["is_current"] is False

    current = rows["Senior Technical Writer"]
    assert current["company"] == "Willow Developer Tools"
    assert current["start_date"] == "2020-08-01"
    assert current["end_date"] is None
    assert current["date_label"] == "August 2020 — Present"
    assert current["is_current"] is True


def test_writer_role_cannot_borrow_an_unrelated_employer_heading():
    data = recall_response()
    data["work_experiences"][3].update(
        company="Quartz Product Studio",
        employer_evidence="Quartz Product Studio",
    )

    parsed = validate_interpretation(data, RECALL_TEXT)

    assert "Technical Writer" not in {
        row["role"] for row in parsed["work_experiences"]
    }
    assert {
        "category": "work_experience.employer_association",
        "item_index": 3,
    } in parsed["_validation_rejections"]


def test_unsupported_dates_and_achievements_are_removed_without_losing_role():
    data = response()
    data["work_experiences"][0].update(
        start="Jan 2010",
        details="Increased revenue 500%",
    )
    data["work_experiences"][0]["achievements"].extend(
        [
            fact("Increased revenue 500%"),
            fact("Shipped a Python and Kafka platform."),
        ]
    )
    parsed = validate_interpretation(data, TEXT)
    row = parsed["work_experiences"][0]
    assert len(parsed["work_experiences"]) == 2
    assert row["start_date"] is None
    assert row["date_label"].startswith("Unknown —")
    assert row["description"] == "Built a Python service used by 12 teams."
    assert {rejection["category"] for rejection in parsed["_validation_rejections"]} >= {
        "work_experience.date_evidence",
        "work_experience.detail_evidence",
        "work_experience.achievement_evidence",
    }


def test_unsupported_current_status_does_not_mark_role_current():
    data = response()
    data["work_experiences"][0]["end"] = "Current"
    parsed = validate_interpretation(data, TEXT)
    row = parsed["work_experiences"][0]
    assert row["is_current"] is False
    assert row["date_label"] == "Jan 2020 — Unknown"
    assert {
        "category": "work_experience.current_status",
        "item_index": 0,
        "field": "end",
    } in parsed["_validation_rejections"]


def test_unknown_end_not_current_and_zero_distinct_from_unknown():
    assert total_experience([]) is None
    assert total_experience([{"start_date": "2020-01-01", "end_date": None}]) is None
    assert total_experience([{"start_date": "2020-01-01", "end_date": "2020-01-01"}]) == 0
    rows = [{"start_date": "2020-01-01", "end_date": "2022-01-01"},
            {"start_date": "2021-01-01", "end_date": "2023-01-01"}]
    assert total_experience(rows) == 3
    data = response()
    data["work_experiences"][1]["end"] = None
    parsed = validate_interpretation(data, TEXT)
    assert not parsed["work_experiences"][1]["is_current"]
    assert "Unknown" in parsed["work_experiences"][1]["date_label"]
    assert parsed["total_experience"] is None


def test_employer_and_dates_cannot_be_borrowed_from_another_role():
    data = response()
    data["work_experiences"][0].update(
        company="Other Studio", start="Jan 2022", end="Present", details="",
        employer_evidence="Manager | Other Studio | Jan 2022 - Present",
        role_evidence="Engineer | Example Studio | Jan 2020 - Dec 2021",
    )
    text = TEXT + "Manager | Other Studio | Jan 2022 - Present"
    parsed = validate_interpretation(data, text)
    assert len(parsed["work_experiences"]) == 1
    assert parsed["work_experiences"][0]["role"] == "Senior Engineer"
    assert parsed["_validation_rejections"] == [
        {
            "category": "work_experience.employer_association",
            "item_index": 0,
        }
    ]


def test_disabled_and_unapproved_never_call_provider(settings, monkeypatch):
    provider = Mock(side_effect=AssertionError("Network must not be used"))
    monkeypatch.setattr("talent.resume_understanding.resume_provider", provider)
    settings.RESUME_AI_ENABLED = False
    assert understand_resume(rules(), Mock())["_interpretation"]["reason"] == "disabled"
    settings.RESUME_AI_ENABLED = True
    settings.RESUME_AI_EXTERNAL_PROCESSING_APPROVED = False
    assert understand_resume(rules(), Mock())["_interpretation"]["source"] == "rules"
    provider.assert_not_called()


def test_mocked_success_and_timeout_fallback(enabled, monkeypatch):
    provider = Mock()
    provider.generate.return_value = AIProviderResult(
        response(), "anthropic", "claude-sonnet-5-5"
    )
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    result = understand_resume(rules(), lambda _tokens: True)
    assert result["_interpretation"]["source"] == "ai"
    assert result["headline"] == "Senior Engineer"
    assert "full_name" not in result  # Local identity is never overridden by AI.
    provider.generate.side_effect = AIProviderTimeout()
    provider.reset_mock()
    result = understand_resume(rules(), lambda _tokens: True)
    assert result["_interpretation"] == {"source": "rules", "reason": "timeout"}
    assert result["headline"] == "Engineer"
    assert provider.generate.call_count == 1


def test_malformed_response_budget_and_no_retry_after_exhaustion(enabled, monkeypatch):
    provider = Mock()
    provider.generate.return_value = AIProviderResult({}, "anthropic", "claude-sonnet-5-5")
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    reserve = Mock(return_value=True)
    record = Mock()
    assert understand_resume(rules(), reserve, record)["_interpretation"]["source"] == "rules"
    assert provider.generate.call_count == reserve.call_count == 2
    assert record.call_count == 2
    assert record.call_args.args[-1] == [
        {"category": "work_experience.schema_validation"}
    ]
    provider.reset_mock()
    understand_resume(rules(), lambda _tokens: False)
    provider.generate.assert_not_called()


def test_reliable_deterministic_history_skips_ai_call(enabled, monkeypatch):
    provider = Mock()
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    data = rules() | {
        "work_experiences": [
            {
                "company": "Example Studio",
                "role": "Engineer",
                "start_date": "2020-01-01",
                "end_date": "2021-01-01",
            }
        ],
        "_employment_parse": {"reliable": True},
    }
    result = understand_resume(data, lambda _tokens: True)
    assert result["_interpretation"] == {
        "source": "rules",
        "reason": "deterministic_reliable",
    }
    provider.generate.assert_not_called()


def test_explicit_candidate_permission_uses_ai_even_when_rules_look_reliable(
    enabled, monkeypatch
):
    provider = Mock()
    provider.generate.return_value = AIProviderResult(
        response(), "anthropic", "claude-sonnet-5-5", stop_reason="end_turn"
    )
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    data = rules() | {
        "work_experiences": [{"company": "Rules Co", "role": "Rules Engineer"}],
        "_employment_parse": {"reliable": True},
    }
    result = understand_resume(
        data,
        lambda _tokens: True,
        candidate_requested_ai=True,
    )
    assert result["_interpretation"]["source"] == "ai"
    provider.generate.assert_called_once()


def test_anthropic_uses_exact_model_structured_output_and_usage(
    enabled, settings, monkeypatch
):
    settings.RESUME_AI_ALLOW_PAID = True
    settings.ANTHROPIC_API_KEY = "synthetic-secret"
    post = Mock(return_value={
        "id": "msg_synthetic",
        "model": "claude-sonnet-5-5",
        "type": "message",
        "content": [
            {"type": "thinking", "thinking": "synthetic reasoning"},
            {"type": "text", "text": "{}"},
        ],
        "stop_reason": "end_turn",
        "usage": {"input_tokens": 321, "output_tokens": 45},
    })
    monkeypatch.setattr("talent.ai_providers._post_json", post)
    result = resume_provider().generate(
        "instruction", {"document_text": "synthetic"}, SCHEMA, "resume"
    )
    url, body, headers, timeout = post.call_args.args
    assert url == "https://api.anthropic.com/v1/messages"
    assert body["model"] == "claude-sonnet-5-5"
    assert body["output_config"]["format"]["type"] == "json_schema"
    assert body["output_config"]["effort"] == "low"
    api_schema = body["output_config"]["format"]["schema"]
    assert "maxLength" not in str(api_schema)
    assert "maxItems" not in str(api_schema)
    assert headers["Authorization"] == "Bearer synthetic-secret"
    assert headers["anthropic-version"] == "2023-06-01"
    assert timeout <= 60
    assert body["max_tokens"] == 8000
    assert result.input_tokens == 321
    assert result.output_tokens == 45
    assert result.stop_reason == "end_turn"


@pytest.mark.parametrize(
    ("payload", "stage"),
    [
        (
            {
                "id": "msg_bad_json",
                "type": "message",
                "model": "claude-sonnet-5-5",
                "content": [{"type": "text", "text": "{not-json"}],
                "stop_reason": "end_turn",
                "usage": {"input_tokens": 300, "output_tokens": 20},
            },
            "json_decode",
        ),
        (
            {
                "id": "msg_truncated",
                "type": "message",
                "model": "claude-sonnet-5-5",
                "content": [{"type": "text", "text": "{\"partial\":"}],
                "stop_reason": "max_tokens",
                "usage": {"input_tokens": 500, "output_tokens": 8000},
            },
            "truncated_output",
        ),
    ],
)
def test_anthropic_malformed_and_truncated_responses_keep_safe_usage_metadata(
    enabled, settings, monkeypatch, payload, stage
):
    settings.RESUME_AI_ALLOW_PAID = True
    settings.ANTHROPIC_API_KEY = "synthetic-secret"
    monkeypatch.setattr("talent.ai_providers._post_json", Mock(return_value=payload))

    with pytest.raises(AIProviderMalformedResponse) as exc_info:
        resume_provider().generate("instruction", {}, SCHEMA, "resume")

    error = exc_info.value
    assert error.stage == stage
    assert error.request_id == payload["id"]
    assert error.input_tokens == payload["usage"]["input_tokens"]
    assert error.output_tokens == payload["usage"]["output_tokens"]
    assert error.stop_reason == payload["stop_reason"]
    assert error.usage_confirmed is True


def test_anthropic_http_error_keeps_request_id_without_provider_message(
    enabled, settings, monkeypatch
):
    settings.RESUME_AI_ALLOW_PAID = True
    settings.ANTHROPIC_API_KEY = "synthetic-secret"
    body = BytesIO(json.dumps({
        "type": "error",
        "error": {"type": "authentication_error", "message": "secret detail"},
        "request_id": "req_synthetic",
    }).encode())

    def rejected(*args, **kwargs):
        raise urllib.error.HTTPError(
            "https://api.anthropic.com/v1/messages",
            401,
            "Unauthorized",
            {"request-id": "req_header"},
            body,
        )

    monkeypatch.setattr("talent.ai_providers.urllib.request.urlopen", rejected)
    with pytest.raises(AIProviderUnavailable) as exc_info:
        resume_provider().generate("instruction", {}, SCHEMA, "resume")
    assert exc_info.value.stage == "http_error"
    assert exc_info.value.request_id == "req_synthetic"
    assert "secret detail" not in str(exc_info.value)


def test_anthropic_resume_requires_explicit_paid_approval(enabled, settings):
    settings.ANTHROPIC_API_KEY = "synthetic-secret"
    settings.RESUME_AI_ALLOW_PAID = False
    with pytest.raises(AIProviderUnavailable):
        resume_provider()


def test_direct_identifiers_redacted_without_claiming_anonymity():
    data = {"full_name": "Sample Person", "email": "sample@example.test",
            "phone": "9999999999", "raw_text": "Sample Person sample@example.test 9999999999"}
    text = interpretation_text(data)
    assert "Sample" not in text and "@" not in text and "9999" not in text


def test_historical_remote_work_never_becomes_future_preference():
    from io import BytesIO

    from .test_candidate_flow import resume_pdf

    result = parse_resume(BytesIO(resume_pdf([
        "Name: Example Person", "Experience: Engineer|Example Studio|2020-01-01|Present|",
        "Worked remotely for a hybrid team with flexible deployment schedules.",
    ])), "example.pdf")
    assert result["work_preferences"] == []
    assert result["notice_period_days"] is None
    assert result["expected_salary_lpa"] is None


@pytest.mark.parametrize("error", [AIProviderRateLimited(), AIProviderUnavailable()])
def test_outage_and_rate_limit_fall_back_without_request_storm(enabled, monkeypatch, error):
    provider = Mock()
    provider.generate.side_effect = error
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    interpreted = understand_resume(rules(), lambda _tokens: True)
    assert interpreted["_interpretation"]["reason"] == error.kind
    assert provider.generate.call_count == 1


def test_openai_requires_explicit_cost_and_zdr_approval(enabled, settings, monkeypatch):
    settings.RESUME_AI_PROVIDER = "openai"
    settings.RESUME_AI_MODEL = "configured-paid-model"
    settings.OPENAI_API_KEY = "synthetic-secret"
    settings.RESUME_AI_OPENAI_ZDR_APPROVED = False
    with pytest.raises(AIProviderUnavailable):
        resume_provider()
    settings.RESUME_AI_ALLOW_PAID = True
    with pytest.raises(AIProviderUnavailable):
        resume_provider()
    settings.RESUME_AI_OPENAI_ZDR_APPROVED = True
    post = Mock(return_value={"output_text": "{}"})
    monkeypatch.setattr("talent.ai_providers._post_json", post)
    resume_provider().generate("instruction", {}, SCHEMA, "resume")
    assert post.call_args.args[1]["store"] is False
    assert post.call_args.args[1]["max_output_tokens"] == 8000


@pytest.mark.parametrize(
    "model", ["gemini-3.1-flash-lite", "gemini-3.5-flash-lite"]
)
def test_gemini_uses_exact_model_structured_output_and_usage(
    enabled, settings, monkeypatch, model
):
    settings.RESUME_AI_PROVIDER = "gemini"
    settings.RESUME_AI_MODEL = model
    settings.RESUME_AI_API_KEY = "synthetic-generic-key"
    settings.RESUME_AI_GEMINI_FREE_TIER_APPROVED = True
    post = Mock(return_value={
        "responseId": "synthetic-response",
        "candidates": [{
            "finishReason": "STOP",
            "content": {"parts": [{"text": "{}"}]},
        }],
        "usageMetadata": {
            "promptTokenCount": 321,
            "candidatesTokenCount": 45,
            "thoughtsTokenCount": 6,
        },
    })
    monkeypatch.setattr("talent.ai_providers._post_json", post)
    result = resume_provider().generate("instruction", {}, SCHEMA, "resume")
    url, body, headers, timeout = post.call_args.args
    assert url == (
        "https://generativelanguage.googleapis.com/v1beta/"
        f"models/{model}:generateContent"
    )
    assert headers["x-goog-api-key"] == "synthetic-generic-key"
    assert body["generationConfig"]["responseMimeType"] == "application/json"
    provider_schema = body["generationConfig"]["responseJsonSchema"]
    assert provider_schema != SCHEMA
    assert "maxLength" not in json.dumps(provider_schema)
    assert "maxItems" not in json.dumps(provider_schema)
    assert "additionalProperties" not in json.dumps(provider_schema)
    assert "maxLength" in json.dumps(SCHEMA)
    assert "maxItems" in json.dumps(SCHEMA)
    assert "additionalProperties" in json.dumps(SCHEMA)
    schema_keywords = set()

    def collect_schema_keywords(node):
        if not isinstance(node, dict):
            return
        for key, value in node.items():
            if key == "properties":
                schema_keywords.add(key)
                for child in value.values():
                    collect_schema_keywords(child)
            else:
                schema_keywords.add(key)
                if key == "items":
                    collect_schema_keywords(value)

    collect_schema_keywords(provider_schema)
    assert schema_keywords <= {"type", "properties", "required", "items"}
    assert body["generationConfig"]["maxOutputTokens"] == 8000
    assert timeout <= 60
    assert result.input_tokens == 321
    assert result.output_tokens == 51
    assert result.stop_reason == "STOP"


@pytest.mark.parametrize(
    ("http_status", "google_status", "error_type", "failure_category"),
    [
        (400, "INVALID_ARGUMENT", AIProviderInvalidRequest, "invalid_request"),
        (401, "UNAUTHENTICATED", AIProviderAuthenticationError, "authentication_error"),
        (403, "PERMISSION_DENIED", AIProviderPermissionDenied, "permission_denied"),
        (404, "NOT_FOUND", AIProviderNotFound, "not_found"),
        (429, "RESOURCE_EXHAUSTED", AIProviderRateLimited, "rate_limited"),
        (503, "UNAVAILABLE", AIProviderUnavailable, "unavailable"),
    ],
)
def test_gemini_http_errors_retain_only_safe_diagnostics(
    enabled,
    settings,
    monkeypatch,
    http_status,
    google_status,
    error_type,
    failure_category,
):
    settings.RESUME_AI_PROVIDER = "gemini"
    settings.RESUME_AI_MODEL = "gemini-3.1-flash-lite"
    settings.RESUME_AI_API_KEY = "synthetic-secret"
    settings.RESUME_AI_GEMINI_FREE_TIER_APPROVED = True
    body = BytesIO(json.dumps({
        "error": {
            "code": http_status,
            "status": google_status,
            "message": "private provider detail containing document data",
        }
    }).encode())

    def rejected(*args, **kwargs):
        raise urllib.error.HTTPError(
            "https://generativelanguage.googleapis.com/v1beta/"
            "models/gemini-3.1-flash-lite:generateContent",
            http_status,
            "provider error",
            {"x-goog-request-id": "safe-request-id", "Retry-After": "7"},
            body,
        )

    monkeypatch.setattr("talent.ai_providers.urllib.request.urlopen", rejected)
    monkeypatch.setattr(
        "talent.ai_providers.time.monotonic", Mock(side_effect=[10.0, 10.125])
    )
    with pytest.raises(error_type) as exc_info:
        resume_provider().generate("instruction", {}, SCHEMA, "resume")

    error = exc_info.value
    assert error.kind == failure_category
    assert error.stage == "http_error"
    assert error.http_status == http_status
    assert error.provider_error_code == str(http_status)
    assert error.provider_error_status == google_status
    assert error.request_id == "safe-request-id"
    assert error.latency_ms == 125
    assert "private provider detail" not in str(error)


def test_gemini_provider_generated_json_and_safe_metadata(enabled, settings, monkeypatch):
    from talent.ai_providers import _HTTPPayload

    settings.RESUME_AI_PROVIDER = "gemini"
    settings.RESUME_AI_MODEL = "gemini-3.1-flash-lite"
    settings.RESUME_AI_API_KEY = "synthetic-secret"
    settings.RESUME_AI_GEMINI_FREE_TIER_APPROVED = True
    payload = _HTTPPayload(
        {
            "responseId": "synthetic-response",
            "candidates": [{
                "finishReason": "STOP",
                "content": {"parts": [{"text": json.dumps(response())}]},
            }],
            "usageMetadata": {"promptTokenCount": 700, "candidatesTokenCount": 300},
        },
        request_id="safe-transport-id",
        http_status=200,
        latency_ms=432,
    )
    monkeypatch.setattr("talent.ai_providers._post_json", Mock(return_value=payload))

    result = resume_provider().generate("instruction", {}, SCHEMA, "resume")

    assert result.data == response()
    assert result.request_id == "synthetic-response"
    assert result.http_status == 200
    assert result.latency_ms == 432
    assert result.input_tokens == 700
    assert result.output_tokens == 300


@pytest.mark.django_db
def test_paid_budget_exhaustion_falls_back_without_call_or_usage(
    enabled, settings, tmp_path, client, candidate_account, monkeypatch,
):
    from .test_candidate_experience import upload
    from .test_candidate_flow import resume_pdf

    settings.MEDIA_ROOT = tmp_path
    settings.RESUME_AI_ALLOW_PAID = True
    settings.RESUME_AI_MODEL = "synthetic-paid-model"
    settings.RESUME_AI_DAILY_BUDGET_USD = Decimal("0.001")
    settings.RESUME_AI_MONTHLY_BUDGET_USD = Decimal("0.001")
    settings.RESUME_AI_INPUT_USD_PER_MILLION = Decimal("100")
    settings.RESUME_AI_OUTPUT_USD_PER_MILLION = Decimal("100")
    provider = Mock()
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    result = upload(
        client,
        resume_pdf(["Name: Asha Rao", "Career details available on request."]),
        "synthetic.pdf",
    ).data
    provider.generate.assert_not_called()
    assert result["work_experiences"] == []
    stored = result["latest_resume"]["id"]
    from talent.models import Resume, ResumeAIUsage

    resume = Resume.objects.get(pk=stored)
    assert resume.extracted_data["_interpretation"]["reason"] == "budget_exhausted"
    assert not ResumeAIUsage.objects.filter(resume=resume).exists()


@pytest.mark.django_db
def test_free_tier_gemini_still_enforces_paid_rate_reservation(
    enabled, settings, tmp_path, client, candidate_account, monkeypatch,
):
    from .test_candidate_experience import upload
    from .test_candidate_flow import resume_pdf

    settings.MEDIA_ROOT = tmp_path
    settings.RESUME_AI_PROVIDER = "gemini"
    settings.RESUME_AI_MODEL = "gemini-3.1-flash-lite"
    settings.RESUME_AI_GEMINI_FREE_TIER_APPROVED = True
    settings.RESUME_AI_DAILY_BUDGET_USD = Decimal("0.001")
    settings.RESUME_AI_MONTHLY_BUDGET_USD = Decimal("0.001")
    settings.RESUME_AI_INPUT_USD_PER_MILLION = Decimal("0.25")
    settings.RESUME_AI_OUTPUT_USD_PER_MILLION = Decimal("1.50")
    provider = Mock()
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)

    result = upload(
        client,
        resume_pdf(["Name: Asha Rao", "Career details available on request."]),
        "synthetic-gemini.pdf",
    ).data

    provider.generate.assert_not_called()
    assert result["work_experiences"] == []
    from talent.models import Resume, ResumeAIUsage

    resume = Resume.objects.get(pk=result["latest_resume"]["id"])
    assert resume.extracted_data["_interpretation"]["reason"] == "budget_exhausted"
    assert not ResumeAIUsage.objects.filter(resume=resume).exists()


@pytest.mark.django_db
def test_free_tier_gemini_records_zero_billed_cost_not_reserved_price(
    settings, candidate_account,
):
    from django.contrib.auth.models import User
    from talent.models import Resume, ResumeAIUsage
    from talent.resume_processing import _record_ai_attempt

    settings.RESUME_AI_PROVIDER = "gemini"
    settings.RESUME_AI_GEMINI_FREE_TIER_APPROVED = True
    settings.RESUME_AI_ALLOW_PAID = False
    settings.RESUME_AI_INPUT_USD_PER_MILLION = Decimal("0.25")
    settings.RESUME_AI_OUTPUT_USD_PER_MILLION = Decimal("1.50")
    profile = User.objects.get(pk=candidate_account["user"]["id"]).candidate_profile
    resume = Resume.objects.create(
        candidate=profile,
        file="quarantine/synthetic-gemini.pdf",
        original_name="synthetic-gemini.pdf",
    )
    usage = ResumeAIUsage.objects.create(
        resume=resume,
        attempt=1,
        provider="gemini",
        model="gemini-3.1-flash-lite",
        outcome=ResumeAIUsage.Outcome.RESERVED,
        reserved_input_tokens=2932,
        estimated_cost_usd=Decimal("0.009733"),
    )
    response = AIProviderResult(
        {},
        "gemini",
        "gemini-3.1-flash-lite",
        request_id="synthetic-response",
        input_tokens=900,
        output_tokens=400,
        stop_reason="STOP",
    )

    _record_ai_attempt(usage.pk, response, None, 125)

    usage.refresh_from_db()
    assert usage.estimated_cost_usd == Decimal("0.009733")
    assert usage.actual_cost_usd == Decimal("0")


@pytest.mark.django_db
@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_worker_persists_mocked_ai_interpretation_and_private_evidence(
    enabled, settings, tmp_path, client, candidate_account, monkeypatch, extension,
):
    from io import BytesIO

    from docx import Document
    from talent.models import Resume, ResumeAIUsage

    from .test_candidate_experience import upload
    from .test_candidate_flow import resume_pdf

    settings.MEDIA_ROOT = tmp_path
    lines = TEXT.replace("Sample Person", "Asha Rao").splitlines()
    if extension == "pdf":
        content = resume_pdf(lines)
    else:
        doc = Document()
        for line in lines:
            doc.add_paragraph(line)
        output = BytesIO()
        doc.save(output)
        content = output.getvalue()
    data = response()
    data["full_name"] = fact(None)
    provider = Mock()
    provider.generate.return_value = AIProviderResult(
        data, "anthropic", "claude-sonnet-5-5"
    )
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    result = upload(client, content, f"fictional.{extension}").data
    assert result["full_name"] == "Asha Rao"
    assert result["headline"] == "Senior Engineer"
    assert len(result["work_experiences"]) == 2
    assert result["career_preferences"]["strongest_skills"][0] == "Python"
    assert result["career_preferences"]["standout_signal"] in TEXT
    assert result["career_preferences"]["likely_directions"] == ["Staff Engineering"]
    stored = Resume.objects.get(pk=result["latest_resume"]["id"])
    assert stored.extracted_data["_ai_requests"] == 1
    assert stored.extracted_data["_interpretation"]["source"] == "ai"
    usage = ResumeAIUsage.objects.get(resume=stored)
    assert usage.outcome == ResumeAIUsage.Outcome.SUCCEEDED
    assert usage.provider == "anthropic"
    assert usage.estimated_cost_usd == 0
    assert usage.actual_cost_usd == 0
    assert "_evidence" in stored.extracted_data
    assert "_evidence" not in str(result) and "raw_text" not in str(result)
    assert "Asha Rao" not in provider.generate.call_args.args[1]["document_text"]
    from talent.resume_processing import process_resume_job

    assert process_resume_job(stored.processing_job.pk) == "duplicate"
    assert ResumeAIUsage.objects.filter(resume=stored).count() == 1


@pytest.mark.django_db
def test_seven_role_ai_history_reaches_profile_api_and_work_experience(
    enabled, settings, tmp_path, client, candidate_account, monkeypatch
):
    from datetime import date

    from django.core.files.uploadedfile import SimpleUploadedFile
    from docx import Document
    from talent.models import Resume, ResumeAIUsage, WorkExperience

    from .test_candidate_flow import complete_resume_upload

    settings.MEDIA_ROOT = tmp_path
    settings.RESUME_AI_MAX_ATTEMPTS = 1
    document = Document()
    document.add_paragraph("Name: Asha Rao")
    for line in PROMOTION_TEXT.splitlines():
        if line.strip():
            document.add_paragraph(line)
    content = BytesIO()
    document.save(content)
    provider = Mock()
    provider.generate.return_value = AIProviderResult(
        promotion_response(),
        "anthropic",
        "claude-sonnet-5-5",
        request_id="req_seven_roles",
        input_tokens=1200,
        output_tokens=900,
        stop_reason="end_turn",
    )
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    # The real Aakash document produces no deterministic roles. Force that same branch
    # for this synthetic, repository-safe document even if the rules parser improves.
    monkeypatch.setattr("talent.resume_understanding._needs_ai", lambda _rules: True)

    response = client.post(
        "/api/v1/candidate/resumes/",
        {
            "file": SimpleUploadedFile(
                "fictional-seven-role-history.docx",
                content.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            )
        },
        format="multipart",
    )
    Resume.objects.filter(pk=response.data["latest_resume"]["id"]).update(
        ai_processing_consent_at=timezone.now(),
        ai_processing_consent_version=settings.LEGAL_RESUME_AI_VERSION,
        ai_processing_consent_sha256=settings.LEGAL_RESUME_AI_SHA256,
    )
    profile = complete_resume_upload(client, response).data

    assert len(profile["work_experiences"]) == 7
    assert sum(row["company"] == "CoCubes.com" for row in profile["work_experiences"]) == 3
    assert sum(row["company"] == "PeopleStrong" for row in profile["work_experiences"]) == 3
    year_only = next(
        row for row in profile["work_experiences"]
        if row["company"] == "CoCubes.com" and row["role"] == "Software Engineer"
    )
    assert year_only["start_date"] is None
    assert year_only["end_date"] is None
    assert year_only["date_label"] == "2012 — 2014"
    current = next(row for row in profile["work_experiences"] if row["is_current"])
    assert current["company"] == "PeopleStrong"
    assert current["role"] == (
        "Associate Vice President — AI & Experience (Head of AI Engineering)"
    )
    assert current["start_date"] == "2024-04-01"
    assert current["end_date"] is None

    stored = Resume.objects.get(pk=profile["latest_resume"]["id"])
    assert WorkExperience.objects.filter(candidate=stored.candidate).count() == 7
    assert WorkExperience.objects.filter(
        candidate=stored.candidate,
        company="Altran (formerly Aricent)",
        role="Software Engineer",
        start_date=date(2010, 1, 1),
        end_date=date(2012, 4, 1),
    ).exists()
    assert all(
        row.source_evidence
        for row in WorkExperience.objects.filter(candidate=stored.candidate)
    )
    usage = ResumeAIUsage.objects.get(resume=stored)
    assert usage.outcome == ResumeAIUsage.Outcome.SUCCEEDED
    assert usage.validation_rejections == []


@pytest.mark.django_db
def test_recalled_researcher_and_writer_roles_reach_profile_without_career_break_job(
    enabled, settings, tmp_path, client, candidate_account, monkeypatch
):
    from django.core.files.uploadedfile import SimpleUploadedFile
    from talent.models import Resume, ResumeAIUsage, WorkExperience

    from .test_candidate_flow import complete_resume_upload, resume_pdf

    settings.MEDIA_ROOT = tmp_path
    settings.RESUME_AI_MAX_ATTEMPTS = 1
    provider = Mock()
    provider.generate.return_value = AIProviderResult(
        recall_response(),
        "gemini",
        "gemini-3.5-flash-lite",
        request_id="req_recall_regression",
        input_tokens=500,
        output_tokens=700,
        stop_reason="STOP",
    )
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    monkeypatch.setattr("talent.resume_understanding._needs_ai", lambda _rules: True)
    content = resume_pdf(["Name: Asha Rao", *RECALL_TEXT.splitlines()])

    response = client.post(
        "/api/v1/candidate/resumes/",
        {
            "file": SimpleUploadedFile(
                "fictional-research-writer-history.pdf",
                content,
                content_type="application/pdf",
            )
        },
        format="multipart",
    )
    Resume.objects.filter(pk=response.data["latest_resume"]["id"]).update(
        ai_processing_consent_at=timezone.now(),
        ai_processing_consent_version=settings.LEGAL_RESUME_AI_VERSION,
        ai_processing_consent_sha256=settings.LEGAL_RESUME_AI_SHA256,
    )
    profile = complete_resume_upload(client, response).data

    roles = {(row["company"], row["role"]): row for row in profile["work_experiences"]}
    assert len(roles) == 4
    assert ("Quartz Product Studio", "UX Researcher") in roles
    assert ("Violet API Works", "Technical Writer") in roles
    assert ("Willow Developer Tools", "Senior Technical Writer") in roles
    assert all(role != "Career break" for _company, role in roles)
    assert roles[("Violet API Works", "Technical Writer")]["start_date"] is None
    assert roles[("Violet API Works", "Technical Writer")]["date_label"] == (
        "Unknown — June 2020"
    )
    assert roles[("Willow Developer Tools", "Senior Technical Writer")][
        "is_current"
    ] is True
    stored = Resume.objects.get(pk=profile["latest_resume"]["id"])
    assert WorkExperience.objects.filter(candidate=stored.candidate).count() == 4
    usage = ResumeAIUsage.objects.get(resume=stored)
    assert {
        (
            rejection["category"],
            rejection.get("item_index"),
            rejection.get("field"),
        )
        for rejection in usage.validation_rejections
        if rejection["category"] != "work_experience.detail_evidence"
    } == {
        (
            "work_experience.non_employment",
            1,
            "role",
        )
    }
    assert {
        "category": "work_experience.non_employment",
        "item_index": 1,
        "field": "role",
    } in usage.validation_rejections


@pytest.mark.django_db
def test_worker_records_evidence_rejection_and_confirmed_cost_separately(
    enabled, settings, tmp_path, client, candidate_account, monkeypatch
):
    from talent.ai_providers import AIProviderResult
    from talent.models import Resume, ResumeAIUsage

    from .test_candidate_experience import upload
    from .test_candidate_flow import resume_pdf

    settings.MEDIA_ROOT = tmp_path
    settings.RESUME_AI_ALLOW_PAID = True
    settings.RESUME_AI_MAX_ATTEMPTS = 1
    settings.RESUME_AI_MAX_INPUT_TOKENS = 5000
    settings.RESUME_AI_MAX_OUTPUT_TOKENS = 8000
    settings.RESUME_AI_DAILY_BUDGET_USD = Decimal("0.10")
    settings.RESUME_AI_MONTHLY_BUDGET_USD = Decimal("0.10")
    settings.RESUME_AI_INPUT_USD_PER_MILLION = Decimal("2")
    settings.RESUME_AI_OUTPUT_USD_PER_MILLION = Decimal("10")
    data = response()
    data["full_name"] = fact(None)
    data["work_experiences"][0]["company"] = "Unsupported Company"
    provider = Mock()
    provider.generate.return_value = AIProviderResult(
        data,
        "anthropic",
        "claude-sonnet-5-5",
        request_id="req_evidence_rejected",
        input_tokens=700,
        output_tokens=500,
        stop_reason="end_turn",
    )
    monkeypatch.setattr("talent.resume_understanding.resume_provider", lambda: provider)
    content = resume_pdf(TEXT.replace("Sample Person", "Asha Rao").splitlines())

    result = upload(client, content, "fictional-evidence-rejection.pdf").data

    stored = Resume.objects.get(pk=result["latest_resume"]["id"])
    usage = ResumeAIUsage.objects.get(resume=stored)
    assert stored.extracted_data["_interpretation"]["source"] == "ai"
    assert stored.extracted_data["_interpretation"]["rejected_fields"] == 1
    assert len(stored.extracted_data["work_experiences"]) == 1
    assert usage.outcome == ResumeAIUsage.Outcome.SUCCEEDED
    assert usage.failure_category == ""
    assert usage.failure_stage == ""
    assert usage.request_id == "req_evidence_rejected"
    assert usage.stop_reason == "end_turn"
    assert usage.input_tokens == 700
    assert usage.output_tokens == 500
    assert usage.reserved_input_tokens > 0
    assert usage.estimated_cost_usd == (
        Decimal(usage.reserved_input_tokens) * Decimal("2")
        + Decimal("8000") * Decimal("10")
    ) / Decimal("1000000")
    assert usage.actual_cost_usd == Decimal("0.006400")
    assert usage.validation_rejections == [
        {
            "category": "work_experience.employer_association",
            "item_index": 0,
            "field": "company",
        }
    ]

    provider.generate.side_effect = AIProviderInvalidRequest(
        stage="http_error",
        request_id="req_http_rejected",
        http_status=400,
        provider_error_code="400",
        provider_error_status="INVALID_ARGUMENT",
        latency_ms=217,
    )
    second = upload(client, content, "fictional-provider-error.pdf").data
    second_usage = ResumeAIUsage.objects.get(
        resume_id=second["latest_resume"]["id"]
    )
    assert second_usage.outcome == ResumeAIUsage.Outcome.FAILED
    assert second_usage.failure_category == "invalid_request"
    assert second_usage.failure_stage == "http_error"
    assert second_usage.request_id == "req_http_rejected"
    assert second_usage.http_status == 400
    assert second_usage.provider_error_code == "400"
    assert second_usage.provider_error_status == "INVALID_ARGUMENT"
    assert second_usage.latency_ms == 217
    assert second_usage.reserved_input_tokens > 0
    assert second_usage.actual_cost_usd is None
