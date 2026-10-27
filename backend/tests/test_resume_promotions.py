"""Fictional regression documents; no personal résumé content is stored here."""

from io import BytesIO

import pytest
from docx import Document
from talent.services import parse_resume

from .test_candidate_flow import resume_pdf


def promotion_lines():
    return [
        "Name: Journey Test Candidate",
        "Location: Pune",
        "Skills: Python, Distributed Systems, Kafka, AWS",
        "PROFESSIONAL EXPERIENCE",
        "Northstar (Example Group) May 2019 - Present",
        "Head of AI Engineering Apr 2023 - Present",
        "Built a reliable workflow service used by 12 teams.",
        "Engineering Manager Jan 2021 - Mar 2023",
        "Improved deployment reliability for the platform team.",
        "Senior Engineer May 2019 - Dec 2020",
        "Earlier Studio Jan 2014 - Apr 2019",
        "VP - Engineering (Jun 2017 - Apr 2019) | Senior Engineer (2016-17) | Engineer (2014-16)",
        "Built the assessment service and led its 8-engineer team.",
        "First Company (formerly Example) Jan 2011 - Dec 2013",
        "Software Engineer",
        "Developed internal tools.",
        "EDUCATION",
        "Example University",
    ]


def mixed_history_lines():
    return [
        "Name: History Candidate",
        "PROFESSIONAL EXPERIENCE",
        "Northstar Labs Jan 2020 - Present",
        "Staff Engineer Jan 2023 - Present",
        "Engineering Lead Jun 2021 - Dec 2022",
        "Engineer Jan 2020 - May 2021",
        "Independent Advisor at Side Studio",
        "Jun 2021 - Jun 2022",
        "Earlier Works 2015 - 2017",
        "Developer",
        "Developer at Undated Studio",
        "EDUCATION",
    ]


def resume_file(lines, extension):
    if extension == "pdf":
        return BytesIO(resume_pdf(lines))
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    content = BytesIO()
    document.save(content)
    return content


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_promotions_and_inline_year_ranges(extension):
    lines = promotion_lines()
    content = resume_file(lines, extension)
    parsed = parse_resume(content, f"resume.{extension}")
    rows = parsed["work_experiences"]
    assert len(rows) == 7
    assert parsed["headline"] == "Head of AI Engineering"
    assert parsed["current_company"] == "Northstar (Example Group)"
    assert parsed["standout_signal"] == "Built a reliable workflow service used by 12 teams."
    assert [row["company"] for row in rows[:3]] == ["Northstar (Example Group)"] * 3
    assert [row["company"] for row in rows[3:6]] == ["Earlier Studio"] * 3
    assert rows[-1]["company"] == "First Company (formerly Example)"
    assert rows[4]["date_label"] == "2016-17"
    assert rows[4]["end_date"] == "2017-01-01"
    assert rows[0]["start_date"] == "2023-04-01"
    assert not any(row["role"].startswith("Built") for row in rows)
    assert parsed["skills"][:4] == ["Python", "Distributed Systems", "Kafka", "AWS"]


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_promotions_overlap_gap_and_incomplete_dates_are_not_lost(extension):
    parsed = parse_resume(resume_file(mixed_history_lines(), extension), f"resume.{extension}")
    rows = parsed["work_experiences"]
    assert [(row["company"], row["role"], row["date_label"]) for row in rows] == [
        ("Northstar Labs", "Staff Engineer", "Jan 2023 - Present"),
        ("Northstar Labs", "Engineering Lead", "Jun 2021 - Dec 2022"),
        ("Northstar Labs", "Engineer", "Jan 2020 - May 2021"),
        ("Side Studio", "Independent Advisor", "Jun 2021 - Jun 2022"),
        ("Earlier Works", "Developer", "2015 - 2017"),
        ("Undated Studio", "Developer", "Dates not specified"),
    ]
    assert rows[3]["start_date"] == "2021-06-01"  # Concurrent with Northstar.
    assert rows[4]["start_date"] == "2015-01-01"  # Sorting anchor; label stays year-only.
    assert rows[5]["start_date"] is None
    assert rows[5]["end_date"] is None


def test_description_years_do_not_become_total_experience():
    parsed = parse_resume(
        BytesIO(
            resume_pdf(
                [
                    "Name: Fictional Person",
                    "PROFESSIONAL EXPERIENCE",
                    "Engineer at Studio",
                    "Jan 2010 - Jan 2020",
                    "Served clients for 3 years.",
                ]
            )
        ),
        "resume.pdf",
    )
    assert parsed["total_experience"] == 10


def test_unknown_history_does_not_become_an_achievement_title():
    parsed = parse_resume(
        BytesIO(
            resume_pdf(
                [
                    "Name: Fictional Person",
                    "Title: Built a platform and led its 10-engineer team.",
                    "WORK EXPERIENCE",
                    "Engineer at Undated Studio",
                ]
            )
        ),
        "resume.pdf",
    )
    assert parsed["headline"] == "Engineer"
    assert parsed["work_experiences"][0]["start_date"] is None


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_undated_employer_followed_by_dated_role_is_preserved(extension):
    parsed = parse_resume(
        resume_file(
            [
                "Name: Fictional Person",
                "PROFESSIONAL EXPERIENCE",
                "Example Systems",
                "Senior Platform Engineer Jan 2022 - Present",
                "Built a reliable internal platform.",
                "Earlier Works",
                "Software Engineer Jun 2019 - Dec 2021",
                "EDUCATION",
            ],
            extension,
        ),
        f"resume.{extension}",
    )

    assert [(row["company"], row["role"]) for row in parsed["work_experiences"]] == [
        ("Example Systems", "Senior Platform Engineer"),
        ("Earlier Works", "Software Engineer"),
    ]


def test_docx_table_heading_and_separate_date_cells_are_preserved():
    document = Document()
    heading = document.add_table(rows=1, cols=4).rows[0].cells
    heading[0].text = "WORK HISTORY"
    row = document.tables[0].add_row().cells
    for cell, value in zip(
        row,
        ["Example Labs", "Senior Engineer", "2020-02", "Present"],
        strict=True,
    ):
        cell.text = value
    row = document.tables[0].add_row().cells
    for cell, value in zip(
        row,
        ["Earlier Studio", "Software Developer", "06/2017", "01-2020"],
        strict=True,
    ):
        cell.text = value
    parsed = parse_resume(BytesIO(_document_bytes(document)), "table.docx")
    assert [(row["company"], row["role"]) for row in parsed["work_experiences"]] == [
        ("Example Labs", "Senior Engineer"),
        ("Earlier Studio", "Software Developer"),
    ]
    assert parsed["work_experiences"][0]["is_current"] is True
    assert parsed["work_experiences"][1]["start_date"] == "2017-06-01"
    assert parsed["work_experiences"][1]["end_date"] == "2020-01-01"


def _document_bytes(document):
    content = BytesIO()
    document.save(content)
    return content.getvalue()


@pytest.mark.parametrize('extension', ['pdf', 'docx'])
@pytest.mark.parametrize('stacked', [True, False])
def test_internship_specialization_and_separate_role_headers(extension, stacked):
    """Fictional analogue of a slash-separated internship and subsequent job."""
    entries = [
        ('Orchid Research and Technologies', 'AI Intern / SDE Intern - AI',
         'Jan 2026 - Jul 2026'),
        ('Cedar Recruitment', 'AI Full-Stack Developer', 'Sep 2026 - Present'),
    ]
    lines = ['Name: Fictional Candidate', 'EXPERIENCE']
    for company, role, dates in entries:
        lines.extend([company, role, dates] if stacked else [' | '.join([company, role, dates])])
        lines.append('Built reliable tools for the platform team.')
    rows = parse_resume(resume_file(lines, extension), f'fictional.{extension}')['work_experiences']
    assert [(row['company'], row['role']) for row in rows] == [entry[:2] for entry in entries]
    assert rows[0]['employment_type'] == 'Internship'
    assert rows[1]['employment_type'] == ''  # Do not guess full-time from its duration.
    assert rows[0]['start_date'] == '2026-01-01'
    assert rows[0]['end_date'] == '2026-07-01'
    assert rows[1]['start_date'] == '2026-09-01' and rows[1]['is_current']
    assert rows[0]['description'] == 'Built reliable tools for the platform team.'
    assert 'Cedar' not in rows[0]['source_evidence']
    assert entries[0][0] in rows[0]['source_evidence']
    assert entries[0][1] in rows[0]['source_evidence']
