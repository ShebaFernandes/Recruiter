from io import BytesIO

import pytest
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from rest_framework.authtoken.models import Token
from talent.models import CandidateProfile, Resume, WorkExperience
from talent.services import parse_resume


def resume_docx(name="Asha Rao", role="Senior Backend Engineer", extra=""):
    document = Document()
    document.add_paragraph(f"Name: {name}")
    document.add_paragraph("Email: asha@example.com")
    document.add_paragraph("+919876543210")
    document.add_paragraph(f"Role: {role}")
    document.add_paragraph("Company: SignalWorks")
    document.add_paragraph("Location: Bengaluru")
    document.add_paragraph("6 years of experience")
    document.add_paragraph("Skills: Python, Django, PostgreSQL, Kafka")
    document.add_paragraph("LinkedIn: https://linkedin.com/in/asharao")
    document.add_paragraph("GitHub: https://github.com/asharao")
    document.add_paragraph("Notice period: 30 days")
    document.add_paragraph("Current salary: 22.5")
    document.add_paragraph("Expected salary: 28")
    document.add_paragraph("Summary: Builds reliable event-driven platforms.")
    document.add_paragraph("Education: B.Tech Computer Science, PES University")
    document.add_paragraph(
        "Experience: Senior Backend Engineer|SignalWorks|2022-01-01|Present|Owned payments APIs"
    )
    document.add_paragraph(
        "Experience: Software Engineer|EarlyStack|2019-01-01|2021-08-01|Built Django services"
    )
    if extra:
        document.add_paragraph(extra)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def resume_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    page[NameObject("/Resources")] = resources
    lines = [
        "Name: Priya Nair",
        "Email: priya@example.com",
        "Role: Backend Engineer",
        "Location: Pune",
        "4 years of experience",
        "Skills: Python, Django, PostgreSQL",
    ]
    commands = ["BT", "/F1 12 Tf", "72 730 Td"]
    for index, line in enumerate(lines):
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        if index:
            commands.append("0 -18 Td")
        commands.append(f"({escaped}) Tj")
    commands.append("ET")
    stream = DecodedStreamObject()
    stream.set_data("\n".join(commands).encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def standard_layout_resume_docx():
    document = Document()
    for line in [
        "Sheba Paul Fernandes",
        "EXPERIENCE",
        "Enter Recruitment    Sep 2026 - Present",
        "AI Full-Stack Developer",
        "- Developing a production recruitment platform.",
        "- Building resume extraction and candidate search workflows.",
        "NewSpace Research and Technologies    Jan 2026 - Jul 2026",
        "AI Intern / SDE Intern - AI | Bengaluru, India",
        "- Built a LangGraph-based product requirements agent.",
        "PROJECTS",
        "Candidate matching system",
    ]:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def test_resume_parser_extracts_company_date_then_role_layout():
    extracted = parse_resume(standard_layout_resume_docx(), "standard-layout.docx")

    assert extracted["headline"] == "AI Full-Stack Developer"
    assert extracted["current_company"] == "Enter Recruitment"
    assert extracted["work_experiences"] == [
        {
            "role": "AI Full-Stack Developer",
            "company": "Enter Recruitment",
            "start_date": "2026-09-01",
            "end_date": None,
            "description": (
                "Developing a production recruitment platform. "
                "Building resume extraction and candidate search workflows."
            ),
        },
        {
            "role": "AI Intern / SDE Intern - AI",
            "company": "NewSpace Research and Technologies",
            "start_date": "2026-01-01",
            "end_date": "2026-07-01",
            "description": "Built a LangGraph-based product requirements agent.",
        },
    ]


@pytest.mark.django_db
def test_candidate_signup_upload_extract_edit_and_login_persists(
    client, candidate_account, tmp_path
):
    with override_settings(MEDIA_ROOT=tmp_path):
        upload = SimpleUploadedFile(
            "asha-resume.docx",
            resume_docx(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response = client.post("/api/v1/candidate/resumes/", {"file": upload}, format="multipart")
        assert response.status_code == 201, response.data
        assert response.data["full_name"] == "Asha Rao"
        assert response.data["headline"] == "Senior Backend Engineer"
        assert response.data["location"] == "Bengaluru"
        assert response.data["skills"] == ["Python", "Django", "Kafka", "PostgreSQL"]
        assert response.data["linkedin_url"] == "https://linkedin.com/in/asharao"
        assert response.data["github_url"] == "https://github.com/asharao"
        assert response.data["latest_resume"]["version"] == 1
        assert len(response.data["work_experiences"]) == 2
        assert response.data["education"] == ["B.Tech Computer Science, PES University"]
        assert response.data["profile_status"] == "draft"
        assert response.data["can_submit"] is False
        assert [item["field"] for item in response.data["missing_fields"]] == [
            "work_preferences",
            "visibility",
            "meaningful_work",
        ]
        assert response.data["profile_completion"] == {
            "percent": 77,
            "remaining": 3,
            "required_missing": 2,
            "recommended_missing": 1,
        }
        assert Resume.objects.count() == 1
        assert WorkExperience.objects.count() == 2

        experiences = response.data["work_experiences"]
        experiences[1]["gap_reason"] = "Completed advanced study and independent research."
        update = client.patch(
            "/api/v1/candidate/profile/",
            {
                "headline": "Staff Backend Engineer",
                "expected_salary_lpa": "30.00",
                "skills": ["Python", "Django", "Kafka", "System Design"],
                "meaningful_work": (
                    "I led a payments platform redesign that cut recovery time by 60%."
                ),
                "visibility": "matching_roles",
                "work_preferences": ["Remote", "Hybrid"],
                "work_experiences": experiences,
            },
            format="json",
        )
        assert update.status_code == 200
        assert update.data["headline"] == "Staff Backend Engineer"
        assert update.data["work_experiences"][1]["gap_reason"] == (
            "Completed advanced study and independent research."
        )
        assert update.data["profile_completion"]["percent"] == 100
        assert update.data["can_submit"] is True
        no_consent = client.post("/api/v1/candidate/profile/submit/", format="json")
        assert no_consent.status_code == 400
        assert "consent" in no_consent.data["detail"].lower()
        submitted = client.post(
            "/api/v1/candidate/profile/submit/", {"consent": True}, format="json"
        )
        assert submitted.status_code == 200
        assert submitted.data["profile_status"] == "submitted"
        assert submitted.data["submitted_at"]
        assert submitted.data["submission_consent_at"]
        assert submitted.data["submission_consent_version"] == "candidate-profile-sharing-v1"
        profile_id = submitted.data["id"]

        edited = client.patch(
            "/api/v1/candidate/profile/",
            {"summary": "Updated after formal submission."},
            format="json",
        )
        assert edited.status_code == 200
        assert edited.data["id"] == profile_id
        assert edited.data["profile_status"] == "submitted"
        assert CandidateProfile.objects.count() == 1

        client.post("/api/v1/auth/logout/")
        client.credentials()
        login = client.post(
            "/api/v1/auth/login/",
            {"email": "asha@example.com", "password": "strong-pass-123"},
            format="json",
        )
        assert login.status_code == 200
        client.credentials(HTTP_AUTHORIZATION=f"Token {login.data['token']}")
        persisted = client.get("/api/v1/candidate/profile/")
        assert persisted.status_code == 200
        assert persisted.data["headline"] == "Staff Backend Engineer"
        assert persisted.data["expected_salary_lpa"] == "30.00"
        assert persisted.data["skills"][-1] == "System Design"
        assert persisted.data["meaningful_work"].startswith("I led a payments")
        assert persisted.data["visibility"] == "matching_roles"
        assert persisted.data["work_preferences"] == ["Remote", "Hybrid"]
        assert persisted.data["missing_fields"] == []
        assert persisted.data["profile_status"] == "submitted"
        assert persisted.data["submitted_at"]
        assert CandidateProfile.objects.get().email == "asha@example.com"


@pytest.mark.django_db
def test_candidate_preferences_reject_unknown_work_mode(client, candidate_account):
    response = client.patch(
        "/api/v1/candidate/profile/",
        {"work_preferences": ["Remote", "Anywhere"]},
        format="json",
    )
    assert response.status_code == 400
    assert response.data["work_preferences"] == ["Choose valid work preferences."]


@pytest.mark.django_db
def test_candidate_cannot_submit_with_required_details_missing(client, candidate_account):
    response = client.post("/api/v1/candidate/profile/submit/", format="json")
    assert response.status_code == 400
    assert response.data["detail"] == "Complete the required profile details before submitting."
    missing = {item["field"] for item in response.data["missing_fields"]}
    assert {"resume", "headline", "location", "visibility"} <= missing


@pytest.mark.django_db
def test_resume_extracts_bare_profile_links_and_work_preferences(
    client, candidate_account, tmp_path
):
    document = Document()
    for line in [
        "Name: Asha Rao",
        "Role: Senior Backend Engineer",
        "Company: SignalWorks",
        "Location: Bengaluru",
        "6 years of experience",
        "Skills: Python, Django",
        "Portfolio: linkedin.com/in/asharao github.com/asharao",
        "Preferred work setup: Remote, Hybrid",
        "Notice period: 30 days",
        "Expected salary: 28",
        "Experience: Senior Backend Engineer|SignalWorks|2022-01-01|Present|Owned APIs",
    ]:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    with override_settings(MEDIA_ROOT=tmp_path):
        upload = SimpleUploadedFile(
            "preferences.docx",
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )
        response = client.post(
            "/api/v1/candidate/resumes/", {"file": upload}, format="multipart"
        )
    assert response.status_code == 201
    assert response.data["linkedin_url"] == "https://linkedin.com/in/asharao"
    assert response.data["github_url"] == "https://github.com/asharao"
    assert response.data["work_preferences"] == ["Remote", "Hybrid"]
    missing = {item["field"] for item in response.data["missing_fields"]}
    assert "linkedin_url" not in missing
    assert "github_url" not in missing
    assert "work_preferences" not in missing
    assert "work_experiences" not in missing


@pytest.mark.django_db
def test_resume_extracts_standard_career_and_education_sections(
    client, candidate_account, tmp_path
):
    document = Document()
    for line in [
        "Name: Maya Sen",
        "Email: maya@example.com",
        "Location: Bengaluru",
        "7 years of experience",
        "Skills: Python, Django, PostgreSQL",
        "Notice period: 30 days",
        "WORK EXPERIENCE",
        "Senior Backend Engineer",
        "SignalWorks",
        "Jan 2022 - Present",
        "Owned payments APIs and Kafka consumers.",
        "Software Engineer",
        "EarlyStack",
        "Jun 2019 - Aug 2021",
        "Built Django services.",
        "EDUCATION",
        "B.Tech in Computer Science",
        "PES University",
        "2015 - 2019",
    ]:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    with override_settings(MEDIA_ROOT=tmp_path):
        upload = SimpleUploadedFile(
            "standard-layout.docx",
            output.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )
        response = client.post(
            "/api/v1/candidate/resumes/", {"file": upload}, format="multipart"
        )
    assert response.status_code == 201, response.data
    assert response.data["headline"] == "Senior Backend Engineer"
    assert response.data["current_company"] == "SignalWorks"
    assert [item["company"] for item in response.data["work_experiences"]] == [
        "EarlyStack",
        "SignalWorks",
    ]
    assert response.data["work_experiences"][0]["start_date"] == "2019-06-01"
    assert response.data["work_experiences"][1]["end_date"] is None
    assert response.data["education"] == [
        "B.Tech in Computer Science · PES University · 2015 - 2019"
    ]
    missing = {item["field"] for item in response.data["missing_fields"]}
    assert "work_experiences" not in missing
    assert "education" not in missing


@pytest.mark.django_db
def test_resume_upload_rejects_unsupported_file(client, candidate_account):
    upload = SimpleUploadedFile("resume.txt", b"not an accepted resume", content_type="text/plain")
    response = client.post("/api/v1/candidate/resumes/", {"file": upload}, format="multipart")
    assert response.status_code == 400
    assert "PDF or DOCX" in response.data["detail"]


@pytest.mark.django_db
def test_real_pdf_is_extracted(client, candidate_account, tmp_path):
    with override_settings(MEDIA_ROOT=tmp_path):
        upload = SimpleUploadedFile(
            "priya-resume.pdf", resume_pdf(), content_type="application/pdf"
        )
        response = client.post("/api/v1/candidate/resumes/", {"file": upload}, format="multipart")
        assert response.status_code == 201, response.data
        assert response.data["full_name"] == "Priya Nair"
        assert response.data["headline"] == "Backend Engineer"
        assert response.data["location"] == "Pune"
        assert response.data["skills"] == ["Python", "Django", "PostgreSQL"]
        missing = {item["field"] for item in response.data["missing_fields"]}
        assert {"linkedin_url", "github_url", "work_experiences", "visibility"} <= missing
        assert not {"headline", "location", "skills"} & missing


@pytest.mark.django_db
def test_candidate_cannot_open_another_candidates_resume(client, candidate_account, tmp_path):
    with override_settings(MEDIA_ROOT=tmp_path):
        own_upload = SimpleUploadedFile(
            "own.docx",
            resume_docx(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        own = client.post("/api/v1/candidate/resumes/", {"file": own_upload}, format="multipart")
        own_resume_id = own.data["latest_resume"]["id"]
        client.post("/api/v1/auth/logout/")
        client.credentials()
        other = client.post(
            "/api/v1/auth/signup/",
            {
                "email": "other@example.com",
                "password": "strong-pass-123",
                "full_name": "Other Person",
                "role": "candidate",
            },
            format="json",
        )
        assert other.status_code == 201
        other_user = User.objects.get(username="other@example.com")
        other_user.is_active = True
        other_user.save(update_fields=["is_active"])
        other_user.candidate_profile.email_verified_at = timezone.now()
        other_user.candidate_profile.save(
            update_fields=["email_verified_at", "profile_updated_at"]
        )
        other_token = Token.objects.create(user=other_user)
        client.credentials(HTTP_AUTHORIZATION=f"Token {other_token.key}")
        response = client.get(f"/api/v1/resumes/{own_resume_id}/download/")
        assert response.status_code == 404
        assert client.get("/api/v1/resumes/99999999/download/").status_code == 404
