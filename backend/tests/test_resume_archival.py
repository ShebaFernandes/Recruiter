from io import BytesIO
from uuid import uuid4

import pytest
from django.utils import timezone
from docx import Document
from talent.employment_reconciliation import (
    complete_employment_history,
    reconcile_employment,
    trusted_rows,
)
from talent.models import CandidateProfile, WorkExperience
from talent.resume_understanding import interpretation_text, validate_interpretation
from talent.services import parse_resume

from .test_candidate_experience import upload
from .test_candidate_flow import resume_pdf

pytestmark = pytest.mark.django_db


def upload_resume(client, content, *, extension='docx'):
    return upload(client, content, name=f'resume-{uuid4().hex}.{extension}')


def structured_resume(*roles):
    document = Document()
    for line in [
        'Name: Asha Rao', 'Role: Developer', 'Work Experience',
        *[f'Work: {title}|{company}|{start}|{end}|Built fictional systems'
          for title, company, start, end in roles],
        'Skills: Python',
    ]:
        document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def standard_resume(*lines):
    document = Document()
    for line in ['Asha Rao', 'EXPERIENCE', *lines, 'SKILLS', 'Python']:
        document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def spaced_pdf_layout_resume(heading, *roles, extra_lines=()):
    document = Document()
    for line in ['Asha Rao', heading]:
        document.add_paragraph(line)
    for company, dates, title, location in roles:
        document.add_paragraph(f'{company}                         {dates}')
        document.add_paragraph(f'{title}                         {location}')
        document.add_paragraph('• Built fictional systems with Python.')
    for line in extra_lines:
        document.add_paragraph(line)
    document.add_paragraph('SKILLS')
    document.add_paragraph('Python')
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_spaced_pdf_layout_one_role_is_complete(client, candidate_account):
    result = upload_resume(client, spaced_pdf_layout_resume(
        'PROFESSIONAL EXPERIENCE',
        ('Harbor Systems', 'Jan 2020 - Present', 'Senior Software Engineer',
         'Bengaluru, India'),
    )).data
    assert [(row['company'], row['role']) for row in result['work_experiences']] == [
        ('Harbor Systems', 'Senior Software Engineer')
    ]
    parsed = CandidateProfile.objects.get().resumes.first().extracted_data
    assert parsed['_employment_parse']['complete_for_archival']
    assert parsed['_employment_parse']['recognized_roles'] == 1
    assert 'Bengaluru, India' not in result['work_experiences'][0]['description']


def test_spaced_pdf_file_extracts_complete_role(client, candidate_account):
    result = upload_resume(client, resume_pdf([
        'Asha Rao', 'PROFESSIONAL EXPERIENCE',
        'Harbor Systems                  Jan 2020 - Present',
        'Senior Software Engineer                  Bengaluru, India',
        'Built fictional systems.', 'SKILLS', 'Python',
    ]), extension='pdf').data
    assert [(row['company'], row['role']) for row in result['work_experiences']] == [
        ('Harbor Systems', 'Senior Software Engineer')
    ]
    assert CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']


def test_spaced_pdf_layout_multiple_roles_promotion_and_internship(client, candidate_account):
    result = upload_resume(client, spaced_pdf_layout_resume(
        'CAREER HISTORY',
        ('Blue Lantern Labs', 'Jan 2018 - Jun 2018', 'Engineering Intern', 'Pune, India'),
        ('Copper Systems', 'Jan 2020 - Jun 2023', 'Software Engineer', 'Pune, India'),
        ('Copper Systems', 'Jul 2023 - Present', 'Senior Software Engineer', 'Pune, India'),
    )).data
    assert len(result['work_experiences']) == 3
    assert {row['role'] for row in result['work_experiences']} == {
        'Engineering Intern', 'Software Engineer', 'Senior Software Engineer'
    }
    assert CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']


def test_spaced_pdf_layout_unknown_role_blocks_archival(client, candidate_account):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    incomplete = upload_resume(client, spaced_pdf_layout_resume(
        'PROFESSIONAL EXPERIENCE',
        ('Blue Lantern Labs', 'Jan 2018 - Jun 2018', 'Engineering Intern', 'Pune, India'),
        extra_lines=('Unpaired Fictional Company', 'Research Associate'),
    )).data
    assert not CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']
    assert {row['id'] for row in incomplete['work_experiences']} == {
        row['id'] for row in first['work_experiences']
    }
    assert incomplete['archived_work_experiences'] == []


def test_selected_experience_extracts_role_without_claiming_complete_history(
    client, candidate_account
):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    selected = upload_resume(client, resume_pdf([
        'SELECTED EXPERIENCE',
        'Engineering Intern | Blue Lantern Labs | Jan 2018 - Jun 2018',
        'SKILLS', 'Python',
    ]), extension='pdf').data
    assert CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['recognized_roles'] == 1
    assert not CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']
    assert {row['id'] for row in selected['work_experiences']} == {
        row['id'] for row in first['work_experiences']
    }
    assert selected['archived_work_experiences'] == []


def test_validated_gemini_role_can_save_without_deterministic_role(client, candidate_account):
    document = Document()
    for line in ('Fictional Candidate', 'Career details', 'Engineer at Fictional Studio'):
        document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    content = stream.getvalue()
    rules = parse_resume(BytesIO(content), 'fictional.docx')
    assert rules['work_experiences'] == []
    null_fact = {'value': None, 'evidence': ''}
    payload = {
        'full_name': null_fact, 'location': null_fact,
        'professional_title': null_fact, 'professional_summary': null_fact,
        'standout_signal': null_fact, 'skills': [], 'education': [],
        'work_experiences': [{
            'company': 'Fictional Studio', 'role': 'Engineer',
            'employer_evidence': 'Engineer at Fictional Studio',
            'role_evidence': 'Engineer at Fictional Studio',
        }],
    }
    validated = validate_interpretation(payload, interpretation_text(rules))
    assert len(trusted_rows(validated)) == 1
    validated['_interpretation'] = {'source': 'ai'}
    validated['_employment_parse'] = rules['_employment_parse']
    assert not complete_employment_history(validated)
    profile = CandidateProfile.objects.get()
    reconcile_employment(profile, trusted_rows(validated))
    assert list(profile.work_experiences.values_list('company', 'role')) == [
        ('Fictional Studio', 'Engineer')
    ]


INTERN = ('Engineering Intern', 'Blue Lantern Labs', '2018-01-01', '2018-06-01')
ENGINEER = ('Developer', 'Copper Systems', '2020-01-01', '2023-06-01')
LEAD = ('Lead Engineer', 'Silver Harbor Systems', '2024-01-01', 'Present')


def test_complete_replacement_archives_omission_and_restores_same_id(client, candidate_account):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    profile = CandidateProfile.objects.get()
    assert first['resume_review'] is None
    assert len(first['work_experiences']) == 2
    assert all(row.origin == 'resume_derived' for row in profile.work_experiences.all())
    assert profile.resumes.first().extracted_data['_employment_parse']['complete_for_archival']
    original = {row['company']: row['id'] for row in first['work_experiences']}

    fewer = upload_resume(client, structured_resume(INTERN)).data
    assert [row['id'] for row in fewer['work_experiences']] == [original[INTERN[1]]]
    assert [row['id'] for row in fewer['archived_work_experiences']] == [original[ENGINEER[1]]]
    assert fewer['resume_review'] is None
    assert fewer['employment_extraction']['profile_state'] == 'applied'
    archived = WorkExperience.all_objects.get(pk=original[ENGINEER[1]])
    assert archived.archived_at is not None
    assert archived.archived_by_resume_id == profile.resumes.first().pk

    repeated = upload_resume(client, structured_resume(INTERN)).data
    assert len(repeated['work_experiences']) == 1
    assert len(repeated['archived_work_experiences']) == 1
    assert WorkExperience.all_objects.filter(candidate=profile).count() == 2

    restored = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    assert {row['id'] for row in restored['work_experiences']} == set(original.values())
    assert restored['archived_work_experiences'] == []
    assert WorkExperience.all_objects.filter(candidate=profile).count() == 2


def test_complete_standard_layout_replacement_archives_omission(client, candidate_account):
    intern = 'Blue Lantern Labs | Engineering Intern | Jan 2018 - Jun 2018'
    engineer = 'Copper Systems | Developer | Jan 2020 - Jun 2023'
    first = upload_resume(client, standard_resume(intern, engineer)).data
    assert len(first['work_experiences']) == 2
    assert CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']
    fewer = upload_resume(client, standard_resume(intern)).data
    assert len(fewer['work_experiences']) == 1
    assert fewer['work_experiences'][0]['id'] == first['work_experiences'][0]['id']
    assert len(fewer['archived_work_experiences']) == 1


def test_unpaired_dated_header_blocks_standard_layout_archiving(client, candidate_account):
    intern = 'Blue Lantern Labs | Engineering Intern | Jan 2018 - Jun 2018'
    engineer = 'Copper Systems | Developer | Jan 2020 - Jun 2023'
    first = upload_resume(client, standard_resume(intern, engineer)).data
    # The second date range lacks a supported role/employer pairing.
    incomplete = upload_resume(client, standard_resume(
        intern, 'Jan 2020 - Jun 2023', 'Undecipherable history line.',
    )).data
    assert not CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']
    assert {row['id'] for row in incomplete['work_experiences']} == {
        row['id'] for row in first['work_experiences']
    }
    assert incomplete['archived_work_experiences'] == []


def test_undated_unpaired_employer_blocks_standard_layout_archiving(client, candidate_account):
    intern = 'Blue Lantern Labs | Engineering Intern | Jan 2018 - Jun 2018'
    engineer = 'Copper Systems | Developer | Jan 2020 - Jun 2023'
    first = upload_resume(client, standard_resume(intern, engineer)).data
    incomplete = upload_resume(client, standard_resume(
        intern, 'Unpaired Fictional Company',
    )).data
    assert not CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']
    assert {row['id'] for row in incomplete['work_experiences']} == {
        row['id'] for row in first['work_experiences']
    }
    assert incomplete['archived_work_experiences'] == []


def test_added_role_and_promotion_keep_existing_id(client, candidate_account):
    first = upload_resume(client, structured_resume(INTERN)).data
    intern_id = first['work_experiences'][0]['id']
    second = upload_resume(client, structured_resume(INTERN, ENGINEER, LEAD)).data
    assert len(second['work_experiences']) == 3
    assert intern_id in {row['id'] for row in second['work_experiences']}
    assert second['archived_work_experiences'] == []
    assert second['resume_review'] is None


def test_candidate_confirmed_omission_is_archived_without_losing_edits(client, candidate_account):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    edited = [dict(row) for row in first['work_experiences']]
    for row in edited:
        if row['company'] == ENGINEER[1]:
            row['description'] = 'Candidate corrected fictional responsibilities.'
    corrected = client.patch('/api/v1/candidate/profile/', {
        'work_experiences': edited,
    }, format='json')
    assert corrected.status_code == 200
    corrected_id = next(row['id'] for row in corrected.data['work_experiences']
                        if row['company'] == ENGINEER[1])
    fewer = upload_resume(client, structured_resume(INTERN)).data
    assert len(fewer['work_experiences']) == 1
    assert [row['id'] for row in fewer['archived_work_experiences']] == [corrected_id]
    archived = WorkExperience.all_objects.get(pk=corrected_id)
    assert archived.confidence_status == 'candidate_confirmed'
    assert archived.description == 'Candidate corrected fictional responsibilities.'
    assert fewer['resume_review'] is None


def test_complete_replacement_archives_unknown_legacy_origin(client, candidate_account):
    profile = CandidateProfile.objects.get()
    legacy = WorkExperience.objects.create(
        candidate=profile, company='Legacy Fictional Works', role='Engineer',
        origin='unknown', confidence_status='supported',
    )
    current = upload_resume(client, structured_resume(INTERN)).data
    assert legacy.pk not in {row['id'] for row in current['work_experiences']}
    assert legacy.pk in {row['id'] for row in current['archived_work_experiences']}
    legacy.refresh_from_db()
    assert legacy.origin == 'unknown'


def test_complete_manual_title_conflict_saves_latest_and_archives_correction(
    client, candidate_account,
):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    edited = [dict(row) for row in first['work_experiences']]
    for row in edited:
        if row['company'] == ENGINEER[1]:
            row['role'] = 'Candidate corrected title'
    corrected = client.patch('/api/v1/candidate/profile/', {
        'work_experiences': edited,
    }, format='json')
    assert corrected.status_code == 200
    corrected_id = next(row['id'] for row in corrected.data['work_experiences']
                        if row['role'] == 'Candidate corrected title')
    latest = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    assert {row['role'] for row in latest['work_experiences']} == {
        INTERN[0], ENGINEER[0],
    }
    assert latest['resume_review'] is None
    assert [row['id'] for row in latest['archived_work_experiences']] == [corrected_id]
    archived = WorkExperience.all_objects.get(pk=corrected_id)
    assert archived.role == 'Candidate corrected title'
    assert archived.confidence_status == 'candidate_confirmed'
    repeated = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    assert len(repeated['work_experiences']) == 2
    assert len(repeated['archived_work_experiences']) == 1


def test_incomplete_or_failed_extraction_never_archives(client, candidate_account):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    original_ids = {row['id'] for row in first['work_experiences']}
    document = Document()
    for line in ['Name: Asha Rao', 'Role: Developer',
                 'Experience: Engineering Intern|Blue Lantern Labs|2018-01-01|2018-06-01']:
        document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    incomplete = upload_resume(client, stream.getvalue()).data
    assert set(row['id'] for row in incomplete['work_experiences']) == original_ids
    assert incomplete['archived_work_experiences'] == []
    assert not CandidateProfile.objects.get().resumes.first().extracted_data[
        '_employment_parse']['complete_for_archival']

    zero = upload_resume(client, structured_resume()).data
    assert set(row['id'] for row in zero['work_experiences']) == original_ids
    assert zero['archived_work_experiences'] == []


def test_incomplete_replacement_keeps_candidate_correction_active(client, candidate_account):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    corrected = client.patch('/api/v1/candidate/profile/', {
        'work_experiences': first['work_experiences'],
    }, format='json')
    assert corrected.status_code == 200
    corrected_ids = {row['id'] for row in corrected.data['work_experiences']}
    incomplete = upload_resume(client, standard_resume(
        'Blue Lantern Labs | Engineering Intern | Jan 2018 - Jun 2018',
        'Unpaired Fictional Company',
    )).data
    assert {row['id'] for row in incomplete['work_experiences']} == corrected_ids
    assert incomplete['archived_work_experiences'] == []


def test_successful_http_or_incomplete_validation_is_not_archive_proof():
    rows = [
        {'company': company, 'role': title, 'start_date': start,
         'end_date': None if end == 'Present' else end, 'is_current': end == 'Present',
         'date_label': f'{start} — {end}', 'source_evidence': 'fictional quote',
         'confidence_status': 'supported'}
        for title, company, start, end in (INTERN, ENGINEER)
    ]
    base = {
        '_employment_parse': {'complete_for_archival': True, 'rule_role_keys': rows},
        '_text_truncated': False,
        '_interpretation': {'source': 'ai'},
        'work_experiences': rows[:1],
        '_validation_rejections': [],
    }
    assert not complete_employment_history(base)
    base['work_experiences'] = rows
    assert complete_employment_history(base)
    base['_text_truncated'] = True
    assert not complete_employment_history(base)
    base['_text_truncated'] = False
    base['_validation_rejections'] = [{'category': 'work_experience.role_evidence'}]
    assert not complete_employment_history(base)
    base['_validation_rejections'] = [{'category': 'schema.overlong_text_discarded',
                                       'count': 1}]
    assert complete_employment_history(base)
    base['_validation_rejections'] = [
        {'category': 'schema.overlong_text_discarded', 'count': 1},
        {'category': 'work_experience.achievement_evidence'},
        {'category': 'work_experience.detail_evidence'},
    ]
    assert complete_employment_history(base)
    base['_validation_rejections'] = [{'category': 'work_experience.date_evidence'}]
    assert not complete_employment_history(base)
    base['_validation_rejections'] = []
    base['_interpretation'] = {'source': 'rules', 'reason': 'malformed_response'}
    assert not complete_employment_history(base)


def test_owner_can_restore_archived_role_as_protected_history(client, candidate_account):
    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    archived_id = next(row['id'] for row in first['work_experiences']
                       if row['company'] == ENGINEER[1])
    upload_resume(client, structured_resume(INTERN))
    response = client.post(f'/api/v1/candidate/work-experiences/{archived_id}/restore/')
    assert response.status_code == 200
    assert archived_id in {row['id'] for row in response.data['work_experiences']}
    assert response.data['archived_work_experiences'] == []
    restored = WorkExperience.objects.get(pk=archived_id)
    assert restored.confidence_status == 'candidate_confirmed'
    assert restored.origin == 'candidate'
    latest = upload_resume(client, structured_resume(INTERN)).data
    assert archived_id not in {row['id'] for row in latest['work_experiences']}
    assert archived_id in {row['id'] for row in latest['archived_work_experiences']}
    assert WorkExperience.all_objects.get(pk=archived_id).confidence_status == (
        'candidate_confirmed'
    )
    readded = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    assert archived_id in {row['id'] for row in readded['work_experiences']}
    assert readded['archived_work_experiences'] == []
    assert WorkExperience.objects.get(pk=archived_id).confidence_status == 'candidate_confirmed'


def test_other_candidate_cannot_restore_archived_role(client, candidate_account):
    from django.contrib.auth.models import User
    from rest_framework.authtoken.models import Token

    first = upload_resume(client, structured_resume(INTERN, ENGINEER)).data
    archived_id = next(row['id'] for row in first['work_experiences']
                       if row['company'] == ENGINEER[1])
    upload_resume(client, structured_resume(INTERN))
    other = client.post('/api/v1/auth/signup/', {
        'email': 'other@example.com', 'password': 'strong-pass-123',
        'full_name': 'Other Person', 'role': 'candidate',
    }, format='json')
    assert other.status_code == 201
    user = User.objects.get(username='other@example.com')
    user.is_active = True
    user.save(update_fields=['is_active'])
    user.candidate_profile.email_verified_at = timezone.now()
    user.candidate_profile.save(update_fields=['email_verified_at', 'profile_updated_at'])
    client.credentials(HTTP_AUTHORIZATION=f'Token {Token.objects.create(user=user).key}')
    response = client.post(f'/api/v1/candidate/work-experiences/{archived_id}/restore/')
    assert response.status_code == 404
    assert WorkExperience.all_objects.get(pk=archived_id).archived_at is not None
