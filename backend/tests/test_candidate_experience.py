from io import BytesIO

import pytest
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from docx import Document
from rest_framework.test import APIClient
from talent.models import CandidateProfile, Resume, ResumeProcessingJob
from talent.resume_processing import process_resume_job
from talent.services import parse_resume

from .test_candidate_flow import complete_resume_upload, resume_docx, resume_pdf

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def private_test_storage(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path


def upload(client, content=None, name="resume.docx"):
    response = client.post(
        "/api/v1/candidate/resumes/",
        {"file": SimpleUploadedFile(name, content or resume_docx())},
        format="multipart",
    )
    if settings.RESUME_AI_ENABLED and response.status_code == 202:
        Resume.objects.filter(pk=response.data["latest_resume"]["id"]).update(
            ai_processing_consent_at=timezone.now(),
            ai_processing_consent_version=settings.LEGAL_RESUME_AI_VERSION,
            ai_processing_consent_sha256=settings.LEGAL_RESUME_AI_SHA256,
        )
    return complete_resume_upload(client, response)


def test_separate_portfolio_and_multiple_work_modes(client, candidate_account):
    url = "/api/v1/candidate/profile/"
    payload = {
        "github_url": "https://github.com/example",
        "career_preferences": {"portfolio_url": "https://example.com/portfolio"},
        "work_preferences": ["Hybrid", "Remote"],
    }
    assert client.patch(url, payload, format="json").status_code == 200
    data = client.get(url).data
    assert data["github_url"] == payload["github_url"]
    assert data["career_preferences"]["portfolio_url"] == "https://example.com/portfolio"
    assert data["work_preferences"] == ["Hybrid", "Remote"]
    for invalid in ["not a URL", "javascript:alert(1)"]:
        assert client.patch(url, {"career_preferences": {"portfolio_url": invalid}},
                            format="json").status_code == 400
    assert client.patch(url, {"career_preferences": {"portfolio_url": ""}},
                        format="json").status_code == 200
    assert client.get(url).data["github_url"] == payload["github_url"]


def test_career_answers_merge_persist_and_validate(client, candidate_account):
    first = {
        "motivators": ["Exceptional team", "Founder"],
        "exclusions": ["Travel"],
        "custom_exclusions": ["Travel"],
        "preferred_locations": ["Pune"],
        "contact_preferences": ["Email"],
        "availability": "Actively looking",
    }
    assert (
        client.patch(
            "/api/v1/candidate/profile/", {"career_preferences": first}, format="json"
        ).status_code
        == 200
    )
    client.patch(
        "/api/v1/candidate/profile/", {"career_preferences": {"exclusions": []}}, format="json"
    )
    data = client.get("/api/v1/candidate/profile/").data["career_preferences"]
    assert data["motivators"] == first["motivators"]
    assert data["custom_exclusions"] == ["Travel"]
    assert data["exclusions"] == []
    for invalid in [
        {"motivators": ["Fabricated"]},
        {"exclusions": "bad"},
        {"contact_preferences": ["Pigeon"]},
        {"unknown": True},
    ]:
        assert (
            client.patch(
                "/api/v1/candidate/profile/", {"career_preferences": invalid}, format="json"
            ).status_code
            == 400
        )


def test_not_looking_availability_preserves_privacy(client, candidate_account):
    data = client.patch(
        "/api/v1/candidate/profile/",
        {
            "visibility": "approved_recruiters",
            "career_preferences": {"availability": "Not looking right now"},
        },
        format="json",
    ).data
    assert data["visibility"] == "not_looking"


def test_custom_motivations_persist_and_remain_selectable(client, candidate_account):
    url = "/api/v1/candidate/profile/"
    response = client.patch(url, {"career_preferences": {
        "custom_motivators": ["Research time", "Research time"],
        "motivators": ["Research time", "Harder problems"],
    }}, format="json")
    assert response.status_code == 200
    assert response.data["career_preferences"]["custom_motivators"] == ["Research time"]
    assert client.patch(url, {"career_preferences": {"motivators": []}},
                        format="json").status_code == 200
    assert client.patch(url, {"career_preferences": {"motivators": ["Research time"]}},
                        format="json").status_code == 200
    assert client.get(url).data["career_preferences"]["motivators"] == ["Research time"]
    for invalid in ["invalid", [""], ["x" * 241]]:
        assert client.patch(url, {"career_preferences": {"custom_motivators": invalid}},
                            format="json").status_code == 400


@pytest.mark.parametrize("visibility", ["approved_recruiters", "matching_roles", "not_looking"])
def test_career_settings_without_visibility_preserves_saved_choice(
    client, candidate_account, visibility
):
    profile = CandidateProfile.objects.get()
    profile.visibility = visibility
    profile.save(update_fields=["visibility"])
    response = client.patch(
        "/api/v1/candidate/profile/",
        {
            "notice_period_days": 0,
            "work_preferences": ["Remote"],
            "career_preferences": {"contact_preferences": ["Email"]},
        },
        format="json",
    )
    assert response.status_code == 200
    assert response.data["visibility"] == visibility
    profile.refresh_from_db()
    assert profile.visibility == visibility


def test_replacement_review_no_data_loss_and_duplicate_safe(client, candidate_account):
    initial = upload(client).data
    corrected_history = initial["work_experiences"]
    corrected_history[0]["role"] = "Candidate corrected role"
    client.patch(
        "/api/v1/candidate/profile/",
        {
            "headline": "Candidate corrected title",
            "work_experiences": corrected_history,
            "meaningful_work": "My own answer",
            "career_preferences": {"exclusions": ["Travel"]},
        },
        format="json",
    )
    updated = upload(client, resume_docx(role="Staff Engineer")).data
    assert updated["headline"] == "Candidate corrected title"
    assert updated["work_experiences"][0]["role"] == "Candidate corrected role"
    assert updated["resume_review"]["fields"]["headline"] == "Staff Engineer"
    rid = updated["latest_resume"]["id"]
    job = ResumeProcessingJob.objects.get(resume_id=rid)
    assert process_resume_job(job.pk, str(job.idempotency_key)) == "duplicate"
    response = client.post(
        f"/api/v1/candidate/resumes/{rid}/review/",
        {
            "fields": ["headline", "work_experiences"],
            "profile_updated_at": updated["profile_updated_at"],
        },
        format="json",
    )
    assert response.status_code == 200, response.data
    assert response.data["headline"] == "Staff Engineer"
    assert response.data["meaningful_work"] == "My own answer"
    assert response.data["career_preferences"]["exclusions"] == ["Travel"]
    assert response.data["resume_review"] is None
    assert len(response.data["work_experiences"]) == 2
    assert (
        client.post(
            f"/api/v1/candidate/resumes/{rid}/review/",
            {
                "fields": ["headline"],
                "profile_updated_at": updated["profile_updated_at"],
            },
            format="json",
        ).status_code
        == 200
    )


def test_stale_review_rejected(client, candidate_account):
    upload(client)
    client.patch("/api/v1/candidate/profile/", {"headline": "First correction"}, format="json")
    updated = upload(client, resume_docx(role="New resume title")).data
    client.patch("/api/v1/candidate/profile/", {"headline": "Later correction"}, format="json")
    response = client.post(
        f"/api/v1/candidate/resumes/{updated['latest_resume']['id']}/review/",
        {
            "fields": ["headline"],
            "profile_updated_at": updated["profile_updated_at"],
        },
        format="json",
    )
    assert response.status_code == 409
    assert CandidateProfile.objects.get().headline == "Later correction"


def test_original_docx_preview_and_authorization(client, candidate_account):
    document = Document()
    document.add_paragraph("Name: Document Owner")
    paragraph = document.add_paragraph()
    paragraph.add_run("Exact original bold text").bold = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table evidence"
    table.cell(0, 1).text = "Original content"
    buffer = BytesIO()
    document.save(buffer)
    data = upload(client, buffer.getvalue()).data
    rid = data["latest_resume"]["id"]
    response = client.get(f"/api/v1/candidate/resumes/{rid}/preview/")
    assert response.status_code == 200
    assert response["Cache-Control"] == "private, no-store"
    assert response.data["blocks"][1]["runs"][0] == {
        "text": "Exact original bold text",
        "bold": True,
        "italic": False,
    }
    assert response.data["blocks"][2]["rows"] == [["Table evidence", "Original content"]]
    anonymous = APIClient()
    assert anonymous.get(f"/api/v1/candidate/resumes/{rid}/preview/").status_code in (401, 403)
    from django.contrib.auth.models import User
    from django.utils import timezone
    from talent.models import UserRole

    other = User.objects.create_user("other-candidate", password="test-password")
    UserRole.objects.create(user=other, role="candidate")
    CandidateProfile.objects.create(user=other, full_name="Other", email_verified_at=timezone.now())
    anonymous.force_authenticate(other)
    assert anonymous.get(f"/api/v1/candidate/resumes/{rid}/preview/").status_code == 404
    assert (
        anonymous.post(f"/api/v1/candidate/resumes/{rid}/review/", {}, format="json").status_code
        == 409
    )
    assert anonymous.get(f"/api/v1/resumes/{rid}/download/").status_code == 404


def test_real_pdf_extraction_and_download(client, candidate_account):
    data = upload(client, resume_pdf(), "real.pdf").data
    assert data["full_name"] == "Priya Nair"
    response = client.get(f"/api/v1/resumes/{data['latest_resume']['id']}/download/")
    assert response.status_code == 200
    assert b"".join(response.streaming_content).startswith(b"%PDF")


def test_pdf_career_dates_are_extracted_and_saved(client, candidate_account):
    pdf = resume_pdf(
        [
            "Name: Real PDF Candidate",
            "Role: Platform Engineer",
            "Location: Pune",
            "5 years of experience",
            "Skills: Python, AWS",
            "Experience: Engineer|PDF Studio|2019-01-01|2021-08-01|Built APIs",
            "Experience: Platform Engineer|PDF Labs|2021-09-01|Present|Built systems",
        ]
    )
    data = upload(client, pdf, "history.pdf").data
    assert [row["company"] for row in data["work_experiences"]] == ["PDF Studio", "PDF Labs"]
    assert data["work_experiences"][0]["start_date"] == "2019-01-01"
    assert data["work_experiences"][1]["end_date"] is None
    assert "Present" in data["work_experiences"][1]["date_label"]


def test_incomplete_overlapping_history_and_corrections(client, candidate_account):
    document = Document()
    for line in [
        "Name: History Example",
        "Work Experience",
        "Engineer at Alpha",
        "Jan 2020 - Dec 2023",
        "Consultant at Beta",
        "Jun 2022 - Present",
        "Developer at Undated Studio",
        "Education",
        "B.Tech Example University",
    ]:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    data = upload(client, buffer.getvalue()).data
    assert len(data["work_experiences"]) == 3
    assert any(
        row["company"] == "Undated Studio" and row["start_date"] is None
        for row in data["work_experiences"]
    )
    rows = data["work_experiences"]
    rows[0]["role"] = "Corrected Engineer"
    response = client.patch("/api/v1/candidate/profile/", {"work_experiences": rows}, format="json")
    assert response.status_code == 200, response.data
    assert any(row["role"] == "Corrected Engineer" for row in response.data["work_experiences"])
    rows[0]["start_date"], rows[0]["end_date"] = "2024-01-01", "2020-01-01"
    assert (
        client.patch(
            "/api/v1/candidate/profile/", {"work_experiences": rows}, format="json"
        ).status_code
        == 400
    )


def test_evidence_only_signal(client, candidate_account):
    data = upload(
        client, resume_docx(extra="Achievement: Built the real payments recovery service.")
    ).data
    assert (
        data["career_preferences"]["standout_signal"] == "Built the real payments recovery service."
    )
    parsed = parse_resume(BytesIO(resume_docx()), "resume.docx")
    assert not parsed["standout_signal"]


def test_preferences_not_exposed_to_recruiter_serializer(client, candidate_account):
    from types import SimpleNamespace

    from django.contrib.auth.models import AnonymousUser
    from talent.serializers import CandidateProfileSerializer

    profile = CandidateProfile.objects.get()
    profile.career_preferences = {"exclusions": ["Private answer"]}
    profile.save()
    data = CandidateProfileSerializer(
        profile, context={"request": SimpleNamespace(user=AnonymousUser())}
    ).data
    assert "career_preferences" not in data
    assert "resume_review" not in data


def test_missing_experience_is_not_proposed_as_zero(client, candidate_account):
    document = Document()
    document.add_paragraph("Name: No Dates Candidate")
    output = BytesIO()
    document.save(output)
    parsed = parse_resume(BytesIO(output.getvalue()), "no-dates.docx")
    assert parsed["total_experience"] is None


def test_preview_rate_limit_and_quarantine(client, candidate_account):
    from unittest.mock import patch

    from talent.candidate_views import ResumePreviewThrottle
    from talent.models import Resume

    data = upload(client).data
    rid = data["latest_resume"]["id"]
    with patch.object(ResumePreviewThrottle, "rate", "1/min"):
        assert client.get(f"/api/v1/candidate/resumes/{rid}/preview/").status_code == 200
        assert client.get(f"/api/v1/candidate/resumes/{rid}/preview/").status_code == 429
    Resume.objects.filter(pk=rid).update(scan_status="quarantined")
    assert client.get(f"/api/v1/candidate/resumes/{rid}/preview/").status_code == 404


def test_processing_motivations_preserve_preferences_and_do_not_grant_consent(
    client, candidate_account
):
    url = "/api/v1/candidate/profile/"
    before = client.get(url).data
    assert before["career_preferences"].get("motivators", []) == []
    client.patch(url, {"career_preferences": {"preferred_locations": ["Pune"]}}, format="json")
    choices = ["Harder problems", "More ownership", "Exceptional people", "Compensation"]
    response = client.patch(url, {"career_preferences": {"motivators": choices}}, format="json")
    assert response.status_code == 200
    saved = client.get(url).data
    assert saved["career_preferences"]["motivators"] == choices
    assert saved["career_preferences"]["preferred_locations"] == ["Pune"]
    assert saved["visibility"] == before["visibility"]
    assert saved["submission_consent_at"] == before["submission_consent_at"]
    assert not Resume.objects.filter(ai_processing_consent_at__isnull=False).exists()
    cleared = client.patch(url, {"career_preferences": {"motivators": []}}, format="json")
    assert cleared.status_code == 200
    assert client.get(url).data["career_preferences"]["motivators"] == []
